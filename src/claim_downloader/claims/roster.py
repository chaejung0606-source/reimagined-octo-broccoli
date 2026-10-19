"""장학금 지급 명단 엑셀 읽기 — 담당자가 달마다 고쳐 쓰는 표(CLAUDE.md 14절 '장학금 지급 명단').

열 위치를 고정하지 않는다. 담당자 양식이라 유형마다 열이 다르다(실측 2026-10-14: (가)형은 연번·이름·학번·은행·계좌·단가·시간·금액,
(나)형 TA 는 강좌명 열이 끼고 연락처·이메일이 더 있다). 그래서 머리글 글자(공백을 뺀 것)로 열을 찾는다. 머리글은 두 줄일 수 있다
(3행 '산출내역' + 4행 '금액(C=A*B)') — 두 줄을 이어 붙여 본다.

(나)형처럼 한 학생이 두 줄(강좌 둘)이면 이름·학번·은행·계좌 칸이 세로 병합돼 둘째 줄이 비어 있다 — 앞 줄 값을 내려 채운다.
마지막 '합계' 행은 건너뛴다. 앱이 쓰는 값은 이름·학번·금액(·강좌명) 뿐이고 계좌는 ERP 가 채운 값과 **대조만** 한다.

이 파일에 실명·학번은 없다. 테스트는 openpyxl 로 가짜 표를 만든다.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

# 머리글 글자(공백 제거 뒤) → 열 뜻. 앞에 있는 것이 우선.
HEADERS = {
    "name": ("이름", "성명", "성명(한글)"),
    "student_no": ("학번", "사번", "학번/사번", "학번(사번)", "직번"),
    "bank": ("은행", "은행명", "금융기관"),
    "account": ("계좌", "계좌번호"),
    "amount": ("금액(C=A*B)", "지급금액", "지급액", "금액", "청구액", "지원금액", "지원액", "합계금액", "지급총액", "총지급액", "총액"),
    "course": ("강좌명", "강좌", "교과목명", "교과목", "과목명", "과목"),
    # 외국인 학생의 국적 — 담당자가 명단에 열을 더하기로 함(2026-10-14). 비어 있으면 내국인으로 본다
    "nation": ("국적", "국가", "국적(국가)", "국가/지역", "국적(국가/지역)"),
    # 소학회 · 동아리 · 팀 — 진행요원비 명단(`연번 소학회명 이름 학번 은행 계좌`)처럼 개별증빙 파일 이름이 학생 이름이 아니라
    # 소학회명으로 돼 있을 때(`BFS_9월 지원금 신청서.pdf`) 파일을 그 학생에게 붙이는 데 쓴다(담당자 실제 시험 2026-10-19)
    "group": ("소학회명", "소학회", "동아리명", "동아리", "팀명", "학회명", "소속팀"),
}
HOME_NATIONS = ("", "한국", "대한민국", "한국인", "내국인", "kr", "kor", "korea", "southkorea", "republicofkorea",
                "해당없음", "없음", "x")
TOTAL_WORDS = ("합계", "총계", "계")
# 합계 줄 — '합계' '총 합계' '합계(원)' '총계' 처럼 적힌 칸. 이름 칸 오른쪽에 있어도 본다(2026-10-14 점검: '합계(원)' 줄이 마지막 학생의
# 둘째 줄로 들어가 합계 금액이 청구될 뻔했다). '소계' 는 중간 줄이라 건너뛰기만 한다.
TOTAL_RE = re.compile(r"^(총)?(합계|총계|계)(\(.*\))?$|합계|총계")
SUBTOTAL_RE = re.compile(r"소계")
# 머리글 윗줄이 이 글자면 묶음 머리 — 아래 줄 글자로 열을 정한다('산출내역' 아래 단가 · 시간 · 금액)
GROUP_HEADERS = ("산출내역", "지급내역", "지급액내역", "세부내역")
# 가로로 묶은 머리 아래의 '합계' 칸은 금액 열이다 — 프로그램참여지원비 명단의 '활동비/여비(교통비 · 일비 · 식비 · 숙박비 · 합계)'(14절)
GROUP_TOTAL_HEADERS = ("합계", "합계(원)", "계")
MAX_HEADER_SCAN = 12          # 머리글은 위쪽 열두 줄 안에 있다


@dataclass
class Person:
    name: str
    student_no: str = ""          # 숫자만. 9자리면 학번(인적정보 탭), 6자리면 사번(참여인력 탭)
    amount: int | None = None
    bank: str = ""
    account: str = ""             # 적힌 그대로(하이픈 포함). 대조는 digits() 로
    course: str = ""              # (나)형 TA 의 강좌명(교과목) — 행 적요에 들어간다
    nation: str = ""              # 국적 열(있으면) — 외국인 학생만 적는다
    group: str = ""               # 소학회명 · 동아리명 · 팀명 열(있으면) — 개별증빙 파일을 이 이름으로도 찾는다
    row: int = 0                  # 엑셀 행 번호 (사람에게 알려 줄 때)
    note: str = ""                # 이 사람 줄의 문제 (학번 없음 등)

    @property
    def account_digits(self) -> str:
        return re.sub(r"\D", "", self.account)

    @property
    def foreign_nation(self) -> str:
        """명단 국적이 외국이면 그 글자, 비었거나 한국이면 ''. '대한민국(KR)' · '-' 같은 칸도 한국으로 본다 — 그런 칸이 외국인으로
        잡혀 내국인 학생마다 국적 팝업과 노란 띠가 떴다(2026-10-14 점검)."""
        core = re.sub(r"\(.*?\)", "", norm(self.nation)).lower()
        if core in HOME_NATIONS or re.fullmatch(r"[-—–·._/]*", core):
            return ""
        return self.nation


@dataclass
class Roster:
    path: str
    title: str = ""               # B2 제목 — 대표적요의 바탕
    kind: str = ""                # 제목의 '(가)형' 같은 유형 글자: '가' '나' …
    people: list[Person] = field(default_factory=list)
    sheet_total: int | None = None    # 표의 '합계' 행에 적힌 금액
    columns: dict = field(default_factory=dict)   # 열 뜻 → 열 번호 (1부터)
    warnings: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return sum(p.amount or 0 for p in self.people)

    def problems(self) -> list[str]:
        """넣기 전에 사람이 봐야 할 것. 비어 있으면 넣어도 된다."""
        out = list(self.warnings)
        if not self.people:
            out.append("명단에 사람이 없습니다")
        if self.sheet_total is not None and self.sheet_total != self.total:
            out.append(f"표의 합계({self.sheet_total:,})와 금액을 더한 값({self.total:,})이 다릅니다")
        out.extend(self.duplicate_problems())
        for p in self.people:
            if p.note:
                out.append(f"{p.row}행 {p.name}: {p.note}")
        return out


    def duplicate_problems(self) -> list[str]:
        """같은 이름 · 같은 학번 줄 — 담당자 규칙(2026-10-14): 학번이 다르면 다른 학생(동명이인)이라 문제없고,
        학번까지 같은데 교과목이 다르면 한 학생이 여러 교과목 TA 를 한 것이라 문제없다. 학번 · 교과목까지 같으면 같은 줄이
        두 번 들어간 것일 수 있어 알린다. 학번이 같은데 이름이 다르면 명단을 잘못 쓴 것일 수 있어 알린다."""
        out: list[str] = []
        seen: dict[tuple[str, str, str], Person] = {}
        names_by_no: dict[str, set[str]] = {}
        for p in self.people:
            if not p.student_no:
                continue                  # 학번 없음은 그 줄의 note 로 따로 알린다
            name = norm(p.name)
            names_by_no.setdefault(p.student_no, set()).add(name)
            key = (name, p.student_no, norm(p.course))
            if key in seen:
                first = seen[key]
                what = f"교과목 '{p.course}' 까지 같은" if p.course else "같은"
                out.append(f"{first.row}행 · {p.row}행 {p.name}(학번 {p.student_no}) — {what} 줄이 둘입니다. 중복인지 확인하세요")
            else:
                seen[key] = p
        for no, names in names_by_no.items():
            if len(names) > 1:
                out.append(f"학번 {no} 인데 이름이 다릅니다: " + " · ".join(sorted(names)))
        return out


def norm(text) -> str:
    """머리글 비교용 — 공백·줄바꿈을 걷어낸다. '이   름' == '이름'."""
    return re.sub(r"\s+", "", str(text if text is not None else ""))


def digits(text) -> str:
    """숫자만. 202416354.0 처럼 실수로 저장된 학번은 정수로 먼저 바꾼다 — 그대로 두면 끝에 0 이 붙는다."""
    if isinstance(text, float) and text.is_integer():
        text = int(text)
    return re.sub(r"\D", "", str(text if text is not None else ""))


def kind_of(title: str) -> str:
    """'2026년도 7월 (가)형 근로장학금 지급내역' → '가'. 없으면 ''."""
    hit = re.search(r"\(([가-힣])\)\s*형", str(title or ""))
    return hit.group(1) if hit else ""


def _to_int(value) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    text = str(value).strip()
    num = re.sub(r"[^\d.\-]", "", text)
    if not num or num in ("-", "."):
        return None
    try:
        return int(round(float(num)))
    except ValueError:
        return None


def _header_key(label: str) -> str | None:
    """두 줄을 이어 붙인 머리글 글자에서 열 뜻을 찾는다. 정확히 같은 것 → 담고 있는 것 순."""
    if not label:
        return None
    for key, words in HEADERS.items():
        if label in words:
            return key
    for key, words in HEADERS.items():
        if key == "amount" and ("단가" in label or "시간" in label):
            continue
        if key == "nation":
            if "국적" in label:          # '국가' 를 담은 다른 머리글(국가장학금 등)과 헷갈리지 않게 담기 비교는 '국적' 만
                return key
            continue
        if any(w in label for w in words):
            return key
    return None


def _spans(ws) -> set[tuple[int, int]]:
    """가로로 여러 열을 묶은 칸이 덮는 (행, 열) — 묶음 머리('산출내역' · '결제자 정보(이름·은행·계좌)' · '활동비/여비') 줄의 칸들."""
    try:
        return {(m.min_row, c) for m in ws.merged_cells.ranges if m.max_col > m.min_col
                for c in range(m.min_col, m.max_col + 1)}
    except Exception:  # noqa: BLE001
        return set()


def _find_header(ws) -> tuple[int, dict]:
    """(머리글 행 번호, 열 뜻 → 열 번호). 이름 열과 학번(또는 금액) 열이 함께 있는 첫 줄을 머리글로 본다.
    개인 지급 명단(재직자 · 외부 참가자)에는 학번이 없을 수 있다 — 그 사람은 입금처를 사람이 고른다(Person.note)."""
    spans = _spans(ws)
    for r in range(1, min(ws.max_row, MAX_HEADER_SCAN) + 1):
        columns: dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            top = norm(ws.cell(r, c).value)
            below = norm(ws.cell(r + 1, c).value) if r < ws.max_row else ""
            if (r, c) in spans and below:
                # 가로 묶음 머리 — 아래 줄 글자가 열 이름이다. '결제자 정보(이름·은행·계좌)' 는 윗줄에 '이름' 이 들어 있어
                # 윗줄로 보면 은행 · 계좌 열을 놓친다. 아래 줄이 자료(한 줄 머리글을 가로로 묶은 칸)면 윗줄로 돌아간다.
                key = _header_key(below) or ("amount" if below in GROUP_TOTAL_HEADERS else None) or _header_key(top)
            else:
                # 윗줄 글자가 먼저. 윗줄이 비었거나(세로 병합 아래 칸) '산출내역' 같은 묶음 머리면 둘째 줄 글자를 본다
                # 윗줄이 머리글이 아닌 글자면(한 줄 머리글의 '비고' 등) 아래 줄은 자료다 — 첫 학생의 비고를 열 이름으로 보던 것을 막는다
                key = _header_key(top) or (_header_key(below) if below and (not top or top in GROUP_HEADERS) else None)
            if key and key not in columns:
                columns[key] = c
        if ("name" in columns and ("student_no" in columns or "amount" in columns)
                and _header_key(norm(ws.cell(r, columns["name"]).value)) == "name"):
            return r, columns
    return 0, {}


def _merged(ws, r: int, c: int) -> bool:
    """그 칸이 세로 병합의 아래 칸인지((나)형 TA: 한 학생이 두 강좌)."""
    try:
        from openpyxl.cell.cell import MergedCell
        return isinstance(ws.cell(r, c), MergedCell)
    except Exception:  # noqa: BLE001
        return False


def _is_header_row(ws, r: int, columns: dict) -> bool:
    """머리글 둘째 줄인지 — 이름 칸이 비어 있고 글자 칸이 있으면."""
    name = ws.cell(r, columns["name"]).value
    if name not in (None, ""):
        return False
    return any(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip() for c in range(1, ws.max_column + 1))


def _title_above(ws, header_row: int) -> str:
    """머리글 위에서 가장 긴 글자 칸 — B2 제목."""
    best = ""
    for r in range(1, header_row):
        for c in range(1, ws.max_column + 1):
            value = ws.cell(r, c).value
            if isinstance(value, str) and len(value.strip()) > len(best):
                best = value.strip()
    return best


def read(path, sheet: str | None = None) -> Roster:
    """명단 엑셀을 읽는다. 못 읽으면 warnings 에 이유를 적고 사람 없는 Roster 를 돌려준다."""
    roster = Roster(path=str(path))
    try:
        import openpyxl
    except ImportError:
        roster.warnings.append("openpyxl 이 없어 엑셀을 읽지 못합니다 — 앱 준비 단계를 다시 돌리세요")
        return roster
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=False)
    except Exception as error:  # noqa: BLE001 — 열려 있는 파일 · 깨진 파일
        roster.warnings.append("엑셀을 열지 못했습니다: " + (str(error).splitlines()[0][:100] if str(error) else "오류"))
        return roster
    try:
        names = [sheet] if sheet else list(wb.sheetnames)
        found = [(name, _find_header(wb[name])) for name in names]
        found = [(name, hit) for name, hit in found if hit[0]]
        if not found:
            roster.warnings.append("머리글(이름 · 학번)이 있는 줄을 찾지 못했습니다. 시트에 명단 표가 있어야 합니다")
            return roster
        # 시트가 여럿이면(프로그램참여지원비 명단: 국외 · 국내) 사람이 있는 첫 시트 — 둘 다 있으면 알린다(폴더 1개 = 청구서 1개)
        filled = [(name, hit) for name, hit in found if _has_people(wb[name], *hit)] or found
        if len(filled) > 1:
            roster.warnings.append("시트 %d개(%s)에 명단이 있어 '%s' 만 읽었습니다 — 다른 시트는 폴더를 나눠 넣으세요"
                                   % (len(filled), ", ".join(n for n, _ in filled), filled[0][0]))
        name, (header_row, columns) = filled[0]
        ws = wb[name]
        roster.columns = columns
        roster.title = _title_above(ws, header_row)
        roster.kind = kind_of(roster.title)
        start = header_row + 1
        if start <= ws.max_row and _is_header_row(ws, start, columns):
            start += 1
        last: Person | None = None
        found_total = False
        for r in range(start, ws.max_row + 1):
            cells = {key: ws.cell(r, c).value for key, c in columns.items()}
            name = str(cells.get("name") or "").strip()
            amount = _to_int(cells.get("amount"))
            texts = [norm(v) for v in (ws.cell(r, c).value for c in range(1, ws.max_column + 1)) if isinstance(v, str)]
            if norm(name) in TOTAL_WORDS or (not name and any(TOTAL_RE.search(x) for x in texts)):
                roster.sheet_total = amount
                found_total = True
                break
            if not name and any(SUBTOTAL_RE.search(x) for x in texts):
                continue                       # 중간 소계 줄
            if not name and amount is None:
                continue                       # 빈 줄
            if not name and not _merged(ws, r, columns["name"]):
                # 이름 칸이 병합된 둘째 줄이 아닌데 금액만 있다 — 앞 학생 줄로 내려받지 않는다(합계 · 메모 줄일 수 있다)
                roster.warnings.append(f"{r}행: 이름이 없는데 금액({amount:,})이 있습니다 — 합계 줄이면 '합계' 라고 적어 주세요")
                continue
            person = Person(name=name, row=r, amount=amount,
                            student_no=digits(cells.get("student_no")),
                            bank=str(cells.get("bank") or "").strip(),
                            account=str(cells.get("account") or "").strip(),
                            course=str(cells.get("course") or "").strip(),
                            nation=str(cells.get("nation") or "").strip(),
                            group=str(cells.get("group") or "").strip())
            if not name and last is not None:
                # 병합된 둘째 줄 — 이름·학번·은행·계좌(·국적)를 앞 줄에서 내려받는다 ((나)형 TA: 한 학생이 두 강좌)
                person.name, person.student_no = last.name, last.student_no
                person.bank, person.account = person.bank or last.bank, person.account or last.account
                person.nation = person.nation or last.nation
                person.group = person.group or last.group
            if not person.name:
                continue
            if not person.student_no:
                person.note = "학번이 없음 — 소득자는 직접 고르세요"
            elif len(person.student_no) not in (6, 9):
                person.note = f"학번 자릿수가 이상함({len(person.student_no)}자리)"
            if person.amount is None or person.amount <= 0:
                person.note = (person.note + " · " if person.note else "") + "금액이 없음"
            roster.people.append(person)
            last = person
        if roster.people and not found_total:
            roster.warnings.append("명단 아래에 '합계' 줄이 없습니다 — 합계와 맞춰 보지 못했습니다. 금액을 한 번 더 확인하세요")
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    if not roster.title:
        roster.warnings.append("표 위에 제목(예: 2026년도 7월 (가)형 근로장학금 지급내역)이 없습니다 — 대표적요를 직접 적으세요")
    return roster


def _has_people(ws, header_row: int, columns: dict) -> bool:
    """머리글 아래에 이름과 금액이 있는 줄이 하나라도 있는지 — 빈 양식 시트를 건너뛰려고."""
    for r in range(header_row + 1, min(ws.max_row, header_row + 60) + 1):
        name = norm(ws.cell(r, columns["name"]).value)
        if name and name not in TOTAL_WORDS and ("amount" not in columns or _to_int(ws.cell(r, columns["amount"]).value)):
            return True
    return False


def summary_lines(roster: Roster) -> list[str]:
    out = [f"명단: {Path(roster.path).name}", f"제목: {roster.title or '—'}" + (f"  (유형 {roster.kind})" if roster.kind else ""),
           f"사람 {len(roster.people)}명 · 합계 {roster.total:,}원"]
    return out
