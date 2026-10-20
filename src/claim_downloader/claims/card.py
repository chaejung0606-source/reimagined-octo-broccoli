"""카드청구 — 결제 기간 폴더를 읽어 비목마다 청구서 하나 · 결제 건마다 한 줄로 나눈다(담당자 설명 2026-10-20, docs/roadmap.md 3절).

폴더 모양(담당자가 이미 쓰는 것 그대로 — `★★지출진행\\2. 카드\\<기간>\\<비목명> MMDD <건명>\\`):

    26-09(09.09. ~ 10.08.)\\                  결제 기간 = 대표적요의 기간. '26.09.09. ~ 10.08.' 꼴도 읽는다
      기타운영경비 0915 사업단 운영회의\\        결제 1건: 비목 · 사용일(MMDD) · 건명. 금액(`45,000원`)은 있으면 쓴다
        0915 영수증.pdf · 2026-09-15 월 운영회의.hwpx · 서명부.pdf
      기타운영경비\\                            비목 이름만 있는 폴더 아래에 결제 건 폴더를 두어도 된다
        0920 사무용품 구입\\ …

결제 건을 ERP 의 미청구 카드내역에서 고르는 열쇠는 **사용일 + 금액**이고, 금액은 담당자 말대로("증빙 폴더 안에 영수증 pdf파일 or
회의록에 첨부된 영수증 or 카드번호 뒷자리와 지급공문 및 회의록을 보고") 폴더 이름 → 영수증 PDF 글자 → 회의록(.hwpx — 회의록 생성의
'금액' · '회의비 사용처' 칸) → 물품 표 엑셀 합계 → 공문 순으로 찾는다. 같은 날 같은 금액이 여럿이면 가맹점 · 카드 뒷자리 · 승인번호로
좁히고, 그래도 여럿이면 사람에게(`pick`).

청구구분: 회의록이 있으면 (17) 회의비, 물품 표가 있으면 (85) 소모성물품, 둘 다 없으면 (90) 기타. 폴더 이름의 `[회의비]` `[소모품]` `[기타]` 가 먼저.

AI 를 쓰지 않는다(pypdf · zipfile · openpyxl). 그림으로만 된 영수증은 못 읽는다 — 그때는 회의록 · 폴더 이름의 금액을 쓴다.
ERP 카드청구 화면 자동화는 실측 뒤(roadmap 3절) — 이 파일은 폴더 읽기 · 검토 · 결제 건 고르기 규칙만 있다. Qt 없음.
"""
from __future__ import annotations

import datetime as dt
import html
import os
import re
import zipfile
from dataclasses import dataclass, field

from ..naming import BIMOK_NM, normnm
from . import notice as notice_mod
from .schema import CATEGORY_CODES, category_code, parse_amount

START_DAY = 9            # 결제 기간은 9일 ~ 다음 달 8일(14절 '카드청구')
REQ_DIV_NAMES = {"17": "회의비", "85": "소모성물품", "90": "기타"}
# 폴더 이름의 청구구분 표시 — `[회의비]` · `(85)` 같은 것
REQ_DIV_TAGS = {"17": ("회의비", "회의", "17"), "85": ("소모성물품", "소모품", "물품", "85"), "90": ("기타", "90")}
TAG_RE = re.compile(r"[\[(（]\s*([^\])）]{1,8}?)\s*[\])）]")
AMOUNT_IN_NAME_RE = re.compile(r"(?<![\d,])(\d{1,3}(?:,\d{3})+|\d+)\s*원")
# 사용일 — '0915' · '09.15' · '9.15'. 뒤에 '원' 이 오면 금액이다
DATE_IN_NAME_RE = re.compile(r"(?<![\d,.])(\d{1,2})\.?(\d{2})(?![\d,]|\s*원)")
CODE_PREFIX_RE = re.compile(r"^\s*(\d{2})\s*(?:_|\.(?=\s*[가-힣])|\s+(?=[가-힣]))\s*")

PERIOD_FULL_RE = re.compile(          # 26-09(09.09. ~ 10.08.)
    r"(?P<yy>\d{4}|\d{2})\s*-\s*(?P<ym>\d{1,2})\s*\(\s*(?P<m1>\d{1,2})\s*\.\s*(?P<d1>\d{1,2})\s*\.?\s*[~∼-]\s*"
    r"(?P<m2>\d{1,2})\s*\.\s*(?P<d2>\d{1,2})\s*\.?\s*\)")
PERIOD_DOTS_RE = re.compile(          # 26.09.09. ~ 10.08. · 26.09.09.~26.10.08. · 2026.09.09 ~ 2026.10.08
    r"(?P<yy>\d{4}|\d{2})\s*\.\s*(?P<m1>\d{1,2})\s*\.\s*(?P<d1>\d{1,2})\s*\.?\s*[~∼]\s*"
    r"(?:(?P<yy2>\d{4}|\d{2})\s*\.\s*)?(?P<m2>\d{1,2})\s*\.\s*(?P<d2>\d{1,2})\s*\.?")
PERIOD_MONTH_RE = re.compile(r"^\s*(?P<yy>\d{4}|\d{2})\s*-\s*(?P<ym>\d{1,2})\s*(?:카드.*)?$")     # '26-09' 만

RECEIPT_NAME_WORDS = ("영수증", "매출전표", "전표", "카드", "결제", "receipt", "invoice")
RECEIPT_TEXT_WORDS = ("승인번호", "카드번호", "가맹점", "결제금액", "승인금액", "매출전표", "영수증")
MINUTES_NAME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}\s+\S+\s+.+\.(hwpx|hwp|docx)$", re.I)
MINUTES_SUFFIXES = (".hwpx", ".hwp", ".docx", ".pdf")
GOODS_NAME_WORDS = ("구매의뢰", "구입의뢰", "물품", "품명", "소모품", "검수", "구입내역")
EXCEL_SUFFIXES = (".xlsx", ".xlsm")
NOT_RECEIPT_WORDS = ("회의록", "서명부", "공문", "계획", "결과", "사진", "구매의뢰")

STRONG_AMOUNT_WORDS = ("승인금액", "결제금액", "합계", "총액", "총금액", "총결제", "청구금액", "받을금액", "결제요청금액")
WEAK_AMOUNT_WORDS = ("금액", "total", "amount")
MONEY_RE = re.compile(r"(?<![\d.])(\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\s*원|(?<![\d.,])(\d{1,3}(?:,\d{3})+)(?![\d,])")
DATE_RE = re.compile(r"(20\d{2})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})")
CARD_NO_RE = re.compile(r"카\s*드\s*번\s*호\s*[:：]?\s*([0-9*Xx×\-\s]{8,30})")
APPROVAL_RE = re.compile(r"승\s*인\s*번\s*호\s*[:：]?\s*(\d{6,12})")
MERCHANT_RE = re.compile(r"(?:가맹점\s*명|가맹점|상\s*호\s*명?|판매자\s*상호)\s*[:：]?\s*([^\n]+)")


# ── 기간 · 이름 ──────────────────────────────────────────────────────
def _year(yy: str) -> int:
    return int(yy) if len(yy) == 4 else 2000 + int(yy)


def _date(year: int, month: int, day: int) -> dt.date | None:
    try:
        return dt.date(year, month, day)
    except ValueError:
        return None


def parse_period(name: str) -> tuple[dt.date, dt.date] | None:
    """기간 폴더 이름 → (시작일, 끝일). '26-09(09.09. ~ 10.08.)' · '26.09.09. ~ 10.08.' · '26-09'(9일 ~ 다음 달 8일). 못 읽으면 None.
    끝 달이 시작 달보다 작으면 해를 넘긴 것('25-12(12.09. ~ 01.08.)')."""
    hit = PERIOD_FULL_RE.search(name) or PERIOD_DOTS_RE.search(name)
    if hit:
        year = _year(hit.group("yy"))
        m1, d1, m2, d2 = (int(hit.group(k)) for k in ("m1", "d1", "m2", "d2"))
        if "ym" in hit.groupdict() and hit.group("ym") and int(hit.group("ym")) != m1 and m1 == 12 and int(hit.group("ym")) == 1:
            year -= 1            # '27-01(12.09. ~ 01.08.)' 처럼 결제 달로 이름을 단 경우
        year2 = _year(hit.group("yy2")) if hit.groupdict().get("yy2") else year + (1 if m2 < m1 else 0)
        start, end = _date(year, m1, d1), _date(year2, m2, d2)
        return (start, end) if start and end and start <= end else None
    hit = PERIOD_MONTH_RE.match(name)
    if hit:
        year, month = _year(hit.group("yy")), int(hit.group("ym"))
        start = _date(year, month, START_DAY)
        end = _date(year + (month == 12), month % 12 + 1, START_DAY - 1)
        return (start, end) if start and end else None
    return None


def period_text(start: dt.date, end: dt.date) -> str:
    """대표적요의 기간 글자 — '26.09.09.~26.10.08.'(지난 청구서 꼴)."""
    return f"{start:%y.%m.%d}.~{end:%y.%m.%d}."


def category_short(code: str) -> str:
    """대표적요 괄호 안의 비목 이름 — 실제 청구서 `09_[그밖의사업운영경비] 카드 결제 건(26.10.0…` · `03_[교육연구프로그램개발운영비] …`
    (ERP 대시보드 알림 실측 2026-10-20). ERP 보기('기타운영경비')가 아니라 보관 폴더식 이름(naming.BIMOK_NM)에서 점 · 띄어쓰기를 뺀 것이다."""
    return re.sub(r"[\s·ㆍ・]", "", BIMOK_NM.get(code, "") or CATEGORY_CODES.get(code, ""))


def main_summary(code: str, start: dt.date | None, end: dt.date | None) -> str:
    if not code or not start or not end:
        return ""
    return f"{code}_[{category_short(code)}] 카드 결제 건({period_text(start, end)})"


def _category_of(text: str, strict: bool = False) -> str:
    """이름 앞부분(비목 이름 · '09_기타운영경비' · '09') → 비목 번호. 모르면 ''.
    strict 면 비목 이름과 똑같아야 한다(폴더 이름 전체가 비목인지 볼 때 — '기타운영경비 사업단 회의' 는 비목 폴더가 아니다)."""
    text = text.strip()
    if not text:
        return ""
    hit = CODE_PREFIX_RE.match(text + " ")
    number = hit.group(1) if hit and hit.group(1) in CATEGORY_CODES else ""
    rest = text[hit.end():].strip() if hit else text
    code = category_code(rest) if rest else ""
    if not code and rest and not strict:
        found = notice_mod.categories_in(rest)
        code = found[0] if len(found) == 1 else ""
    if number:
        if code and code != number:
            return ""                    # '03_기타운영경비' 처럼 번호와 이름이 다르면 못 정함
        return number if (code or not rest or not strict) else ""
    return code


def is_category_folder(name: str) -> bool:
    """비목 이름만 있는 폴더('기타운영경비' · '09_기타운영경비') — 그 아래가 결제 건 폴더다."""
    return not _date_token(name) and bool(_category_of(name, strict=True))


@dataclass
class ItemName:
    code: str = ""
    used: dt.date | None = None
    title: str = ""
    amount: int | None = None
    req_div: str = ""
    head: str = ""          # 사용일 앞 글자(비목 자리) — 비목을 못 읽었을 때 알림에 쓴다
    mmdd: tuple[int, int] | None = None     # 이름의 사용일(월, 일) — 기간을 못 읽어도 결제 건 폴더인지는 이것으로


def parse_item_name(name: str, period: tuple[dt.date, dt.date] | None = None) -> ItemName:
    """결제 건 폴더 이름 '<비목명> MMDD <건명> [금액원] [회의비]' 을 나눈다. 사용일 해는 기간에서 정한다."""
    out = ItemName()
    rest = name
    for hit in list(TAG_RE.finditer(rest)):
        tag = normnm(hit.group(1))
        for code, words in REQ_DIV_TAGS.items():
            if tag in words:
                out.req_div = out.req_div or code
                rest = rest.replace(hit.group(0), " ")
                break
    amount = AMOUNT_IN_NAME_RE.search(rest)
    if amount:
        out.amount = parse_amount(amount.group(1))
        rest = rest[:amount.start()] + " " + rest[amount.end():]
    date_hit = _date_token(rest)
    if date_hit:
        out.mmdd = (int(date_hit.group(1)), int(date_hit.group(2)))
        out.head = rest[:date_hit.start()].strip(" _-")
        out.code = _category_of(out.head)
        out.used = _use_date(int(date_hit.group(1)), int(date_hit.group(2)), period)
        out.title = re.sub(r"\s+", " ", rest[date_hit.end():]).strip(" _-")
    else:
        out.title = re.sub(r"\s+", " ", rest).strip(" _-")
    return out


def _date_token(text: str):
    for hit in DATE_IN_NAME_RE.finditer(text):
        if 1 <= int(hit.group(1)) <= 12 and 1 <= int(hit.group(2)) <= 31:
            return hit
    return None


def _use_date(month: int, day: int, period: tuple[dt.date, dt.date] | None) -> dt.date | None:
    if not period:
        return None
    start, end = period
    for year in sorted({start.year, end.year}):
        found = _date(year, month, day)
        if found and start <= found <= end:
            return found
    # 기간 밖 — 시작 달 이후면 시작 해, 아니면 끝 해(사람에게 알린다)
    return _date(start.year if month >= start.month else end.year, month, day)


# ── 서류 읽기 ────────────────────────────────────────────────────────
def pdf_text(path: str, pages: int = 3) -> str:
    """PDF 글자(pypdf). 스캔 PDF · 깨진 PDF 는 ''."""
    try:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:pages])
    except Exception:  # noqa: BLE001 — 못 읽는 PDF 는 빈 글
        return ""


@dataclass
class Receipt:
    path: str = ""
    amount: int | None = None
    card_last: str = ""
    approval: str = ""
    merchant: str = ""
    used: dt.date | None = None


def _money_in(line: str) -> list[int]:
    out = []
    for hit in MONEY_RE.finditer(line):
        value = parse_amount(hit.group(1) or hit.group(2))
        if value:
            out.append(value)
    return out


def read_receipt(text: str, path: str = "") -> Receipt:
    """카드 영수증 · 매출전표 글자에서 금액(합계 · 승인금액 줄의 가장 큰 값) · 카드 뒷자리 · 승인번호 · 가맹점 · 날짜."""
    out = Receipt(path=path)
    lines = [line.strip() for line in text.splitlines()]
    strong: list[int] = []
    weak: list[int] = []
    for i, line in enumerate(lines):
        flat = line.replace(" ", "").lower()
        bucket = strong if any(w in flat for w in STRONG_AMOUNT_WORDS) else weak if any(w in flat for w in WEAK_AMOUNT_WORDS) else None
        if bucket is None or any(w in flat for w in ("공급가", "부가세", "세액", "봉사료", "면세")) and bucket is weak:
            continue
        values = _money_in(line)
        if not values and i + 1 < len(lines):
            values = _money_in(lines[i + 1])         # 글자 뽑기에서 값이 다음 줄로 떨어진 PDF
        bucket.extend(values)
    if strong or weak:
        out.amount = max(strong or weak)
    hit = CARD_NO_RE.search(text)
    if hit:
        tail = re.sub(r"[\s\-]", "", hit.group(1))
        if re.search(r"\d{4}$", tail):
            out.card_last = tail[-4:]
    hit = APPROVAL_RE.search(text)
    if hit:
        out.approval = hit.group(1)
    hit = MERCHANT_RE.search(text)
    if hit:
        out.merchant = re.split(r"\s{2,}|사업자|대표자|전화|주소", hit.group(1).strip())[0].strip(" :：")
    for line in lines:
        if any(w in line.replace(" ", "") for w in ("거래일", "승인일", "결제일", "일시", "일자", "이용일")):
            found = DATE_RE.search(line)
            if found:
                out.used = _date(*(int(x) for x in found.groups()))
                break
    if out.used is None:
        found = DATE_RE.search(text)
        out.used = _date(*(int(x) for x in found.groups())) if found else None
    return out


def looks_like_receipt(text: str) -> bool:
    flat = text.replace(" ", "")
    return sum(1 for w in RECEIPT_TEXT_WORDS if w in flat) >= 1


@dataclass
class Minutes:
    path: str = ""
    amount: int | None = None
    merchant: str = ""
    used: dt.date | None = None


MINUTES_LABELS = ("사 업 명", "회의안건", "회의일시", "회의장소", "회의비 사용처", "금액", "참석자", "소속", "이름")


def minutes_texts(path: str) -> list[str]:
    """회의록(.hwpx · .docx) 글자 조각 — 표 칸마다 하나. 회의록 생성(13절)이 만든 꼴. 못 읽으면 []."""
    lower = path.lower()
    member, tag = ("Contents/section0.xml", "hp:t") if lower.endswith(".hwpx") else \
        ("word/document.xml", "w:t") if lower.endswith(".docx") else ("", "")
    if not member:
        return []
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read(member).decode("utf-8", "replace")
    except (OSError, KeyError, zipfile.BadZipFile):
        return []
    return [html.unescape(t).strip() for t in re.findall(rf"<{tag}(?:\s[^>]*)?>(.*?)</{tag}>", xml, re.S)]


def read_minutes(path: str) -> Minutes:
    """회의록의 '금액' · '회의비 사용처' · '회의일시' 칸 — 라벨 바로 다음 글자가 값(비어 있으면 다음 라벨이 온다)."""
    out = Minutes(path=path)
    texts = [t for t in minutes_texts(path) if t]

    def value(label: str) -> str:
        for i, text in enumerate(texts):
            if text.replace(" ", "") == label.replace(" ", "") and i + 1 < len(texts):
                nxt = texts[i + 1]
                return "" if nxt in MINUTES_LABELS else nxt
        return ""

    out.amount = parse_amount(value("금액")) if value("금액") else None
    out.merchant = value("회의비 사용처")
    found = DATE_RE.search(value("회의일시"))
    out.used = _date(*(int(x) for x in found.groups())) if found else None
    return out


@dataclass
class Goods:
    path: str = ""
    sheet: str = ""
    rows: list[dict] = field(default_factory=list)      # 품명 · 규격 · 수량 · 단가 · 금액
    total: int | None = None
    notes: list[str] = field(default_factory=list)


GOODS_HEADERS = {"name": ("품명", "물품명", "품목"), "spec": ("규격",), "qty": ("수량",), "price": ("단가",),
                 "amount": ("금액", "총구입액", "합계금액", "구입금액")}


def read_goods(path: str) -> Goods:
    """물품 표(구매의뢰 엑셀 — 머리글에 '품명') → 줄들 · 합계. 시트가 여럿이면 줄이 있는 첫 시트(여럿이면 알림)."""
    out = Goods(path=path)
    try:
        import openpyxl
        book = openpyxl.load_workbook(path, data_only=True)
    except Exception:  # noqa: BLE001
        out.notes.append(f"{os.path.basename(path)} 을 열지 못함")
        return out
    found: list[tuple[str, list[dict], int | None]] = []
    for ws in book.worksheets:
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 300), values_only=True)]
        sheet_rows, total = _goods_rows(rows)
        if sheet_rows:
            found.append((ws.title, sheet_rows, total))
    if not found:
        out.notes.append(f"{os.path.basename(path)} 에서 '품명' 머리글 아래 물품 줄을 찾지 못함")
        return out
    if len(found) > 1:
        out.notes.append(f"물품 표 시트가 여럿({', '.join(f[0] for f in found)}) — '{found[0][0]}' 을 봄")
    out.sheet, out.rows, total = found[0]
    amounts = [r["amount"] for r in out.rows if isinstance(r.get("amount"), int)]
    out.total = total if total is not None else (sum(amounts) if amounts else None)
    return out


def _goods_rows(rows: list[list]) -> tuple[list[dict], int | None]:
    header_at, columns = -1, {}
    for r, row in enumerate(rows[:25]):
        labels = [normnm(c) for c in row]
        cols = {}
        for key, words in GOODS_HEADERS.items():
            for c, label in enumerate(labels):
                if label and any(label == w or (key != "amount" and label.startswith(w)) or (key == "amount" and w in label
                                                                                         and "공급" not in label and "부가" not in label)
                                 for w in words):
                    cols.setdefault(key, c)
        if "name" in cols:
            header_at, columns = r, cols
            break
    if header_at < 0:
        return [], None
    out, total = [], None
    for row in rows[header_at + 1:]:
        cells = [str(c).strip() if c is not None else "" for c in row]
        flat = "".join(normnm(c) for c in cells)
        if any(w in flat for w in ("합계", "총계", "총액")):
            amount = columns.get("amount")
            if amount is not None and amount < len(row):
                total = parse_amount(row[amount])
            break
        name = cells[columns["name"]] if columns["name"] < len(cells) else ""
        if not name:
            continue
        item = {"name": name}
        for key in ("spec", "qty", "price", "amount"):
            c = columns.get(key)
            value = row[c] if c is not None and c < len(row) else None
            item[key] = (str(value).strip() if value is not None else "") if key == "spec" else parse_amount(value)
        out.append(item)
    return out, total


# ── 결제 건 · 청구서 ──────────────────────────────────────────────────
@dataclass
class CardItem:
    folder: str
    name: str = ""                 # 폴더 이름
    code: str = ""                 # 비목 번호
    used: dt.date | None = None    # 사용일
    title: str = ""                # 건명 — (90) 줄의 적요가 된다
    amount: int | None = None
    amount_from: str = ""          # 폴더 이름 · 영수증 · 회의록 · 물품 표 · 공문
    merchant: str = ""
    card_last: str = ""
    approval: str = ""
    req_div: str = ""              # 17 · 85 · 90
    req_div_from: str = ""
    minutes: str = ""              # 회의록 파일
    goods: Goods | None = None
    receipts: list[str] = field(default_factory=list)
    files: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)       # 확인할 것

    @property
    def req_div_name(self) -> str:
        return f"({self.req_div}) {REQ_DIV_NAMES.get(self.req_div, '')}" if self.req_div else ""

    @property
    def summary(self) -> str:
        """줄 적요 — `NN_<건명>`. (17) · (85) 는 ERP 가 증빙에서 채우는지 미실측(여비는 채웠다 — 14절 끝)."""
        return f"{self.code}_{self.title}" if self.code and self.title else self.title

    def blocking(self) -> list[str]:
        """넣기를 막을 것 — 비목 · 사용일 · 금액 중 하나라도 모르면 ERP 목록에서 결제 건을 고를 수 없다."""
        out = []
        if not self.code:
            out.append("비목을 모름")
        if not self.used:
            out.append("사용일을 모름")
        if self.amount is None:
            out.append("금액을 모름")
        return out


@dataclass
class CardSheet:
    code: str
    start: dt.date | None = None
    end: dt.date | None = None
    items: list[CardItem] = field(default_factory=list)

    @property
    def category(self) -> str:
        return CATEGORY_CODES.get(self.code, "")

    @property
    def main_summary(self) -> str:
        return main_summary(self.code, self.start, self.end)

    @property
    def total(self) -> int:
        return sum(i.amount or 0 for i in self.items)

    def counts(self) -> str:
        by = {code: sum(1 for i in self.items if i.req_div == code) for code in REQ_DIV_NAMES}
        return " · ".join(f"{REQ_DIV_NAMES[c]} {n}" for c, n in by.items() if n)


@dataclass
class CardJob:
    folder: str
    start: dt.date | None = None
    end: dt.date | None = None
    sheets: list[CardSheet] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)    # 결제 건으로 못 읽은 폴더 · 파일
    notes: list[str] = field(default_factory=list)

    @property
    def items(self) -> list[CardItem]:
        return [i for s in self.sheets for i in s.items]

    @property
    def total(self) -> int:
        return sum(s.total for s in self.sheets)

    def problems(self) -> list[str]:
        out = []
        if not self.start:
            out.append("기간 폴더 이름에서 결제 기간을 읽지 못함 — 이름을 '26-09(09.09. ~ 10.08.)' 처럼")
        if not self.items:
            out.append("결제 건 폴더가 없음 — '<비목명> MMDD <건명>' 폴더를 기간 폴더 안에")
        for item in self.items:
            missing = item.blocking()
            if missing:
                out.append(f"{item.name}: " + " · ".join(missing))
        return out

    def summary_lines(self) -> list[str]:
        if not self.start:
            return ["기간: 읽지 못함"]
        out = [f"기간 {period_text(self.start, self.end)} · 청구서 {len(self.sheets)}개 · 결제 {len(self.items)}건 · {self.total:,}원"]
        for sheet in self.sheets:
            out.append(f"{sheet.main_summary or '(비목 모름)'} — 결제 {len(sheet.items)}건 · {sheet.total:,}원"
                       + (f" ({sheet.counts()})" if sheet.counts() else ""))
        return out


def _files_under(folder: str) -> list[str]:
    out = []
    for root, dirs, names in os.walk(folder):
        dirs[:] = sorted(d for d in dirs if not d.startswith("."))
        out += [os.path.join(root, n) for n in sorted(names) if not n.startswith(("~$", ".")) and n != "Thumbs.db"]
    return out


def _is_minutes(path: str) -> bool:
    name = os.path.basename(path)
    return name.lower().endswith(MINUTES_SUFFIXES) and ("회의록" in name or bool(MINUTES_NAME_RE.match(name)))


def _is_goods(path: str) -> bool:
    name = os.path.basename(path)
    return name.lower().endswith(EXCEL_SUFFIXES) and any(w in name.replace(" ", "") for w in GOODS_NAME_WORDS)


def read_item(folder: str, period: tuple[dt.date, dt.date] | None, code_hint: str = "", read_text=pdf_text) -> CardItem:
    """결제 건 폴더 하나 — 이름 · 영수증 · 회의록 · 물품 표 · 공문을 읽고 금액 출처 · 청구구분을 정한다."""
    name = os.path.basename(os.path.normpath(folder))
    parsed = parse_item_name(name, period)
    item = CardItem(folder=folder, name=name, code=parsed.code or code_hint, used=parsed.used, title=parsed.title)
    if parsed.head and not parsed.code and not code_hint:
        item.notes.append(f"'{parsed.head}' 을 비목으로 읽지 못함")
    if parsed.used and period and not (period[0] <= parsed.used <= period[1]):
        item.notes.append(f"사용일 {parsed.used:%m-%d} 이 결제 기간 밖")
    item.files = _files_under(folder)
    pdfs = [p for p in item.files if p.lower().endswith(".pdf")]

    # 영수증 — 이름이 영수증 같은 PDF 먼저, 없으면 글자가 영수증 같은 PDF(회의록 · 공문 · 서명부 빼고)
    named = [p for p in pdfs if any(w in os.path.basename(p).lower() for w in RECEIPT_NAME_WORDS)]
    others = [p for p in pdfs if p not in named and not any(w in os.path.basename(p) for w in NOT_RECEIPT_WORDS)]
    receipts: list[Receipt] = []
    for path in (named + others)[:8]:
        text = read_text(path)
        if text.strip() and (path in named or looks_like_receipt(text)):
            got = read_receipt(text, path)
            if got.amount or got.card_last or got.approval:
                receipts.append(got)
                item.receipts.append(path)
    scanned = [os.path.basename(p) for p in named if p not in item.receipts]
    pictures = [os.path.basename(p) for p in item.files if p.lower().endswith((".jpg", ".jpeg", ".png", ".heic"))
                and any(w in os.path.basename(p).lower() for w in RECEIPT_NAME_WORDS)]
    if scanned or pictures:
        item.notes.append("글자를 못 읽는 영수증: " + ", ".join(scanned + pictures))

    minutes_files = [p for p in item.files if _is_minutes(p)]
    minutes = next((m for m in (read_minutes(p) for p in minutes_files if p.lower().endswith((".hwpx", ".docx")))
                    if m.amount or m.merchant), None)
    if minutes_files:
        item.minutes = minutes.path if minutes else minutes_files[0]
    goods_files = [p for p in item.files if _is_goods(p)]
    if goods_files:
        item.goods = read_goods(goods_files[0])
        item.notes += item.goods.notes
        if not item.goods.rows:
            item.goods = None
    notice = None
    if not item.code or not (parsed.amount or receipts or minutes or item.goods):
        candidates = [p for p in pdfs if p not in item.receipts]
        notice = notice_mod.find(candidates, read_texts=lambda ps: [read_text(p) for p in ps]) if candidates else None
        if notice is not None and not notice.path:
            notice = None
    if not item.code and notice is not None and notice.code:
        item.code = notice.code
        item.notes.append(f"비목은 공문에서: {notice.category}")

    # 금액 — 폴더 이름(사람이 적은 것) → 영수증 → 회의록 → 물품 표 → 공문(계획 금액일 수 있어 약함)
    receipt_amounts = sorted({r.amount for r in receipts if r.amount})
    if len(receipt_amounts) > 1:
        item.notes.append("영수증 금액이 여럿: " + ", ".join(f"{a:,}" for a in receipt_amounts))
    sources = [("폴더 이름", parsed.amount), ("영수증", receipt_amounts[0] if len(receipt_amounts) == 1 else None),
               ("회의록", minutes.amount if minutes else None), ("물품 표", item.goods.total if item.goods else None)]
    known = [(where, value) for where, value in sources if value]
    if known:
        item.amount_from, item.amount = known[0]
        differ = {value for _w, value in known}
        if len(differ) > 1:
            item.notes.append("금액이 서로 다름: " + " · ".join(f"{w} {v:,}" for w, v in known))
    elif notice is not None and notice.amount:
        item.amount, item.amount_from = notice.amount, "공문"
        item.notes.append("금액을 공문에서 읽음(계획 금액일 수 있음) — 영수증 금액과 같은지 확인")

    item.merchant = next((r.merchant for r in receipts if r.merchant), "") or (minutes.merchant if minutes else "")
    item.card_last = next((r.card_last for r in receipts if r.card_last), "")
    item.approval = next((r.approval for r in receipts if r.approval), "")
    other_day = next((r.used for r in receipts if r.used), None) or (minutes.used if minutes else None)
    if item.used is None and other_day:
        item.used = other_day
        item.notes.append(f"사용일은 {'영수증' if receipts else '회의록'}에서: {other_day:%m-%d}")
    elif item.used and other_day and other_day != item.used:
        item.notes.append(f"사용일이 다름: 폴더 {item.used:%m-%d} · {'영수증' if receipts else '회의록'} {other_day:%m-%d}")

    # 청구구분 — 이름 표시 → 회의록 → 물품 표 → 기타
    if parsed.req_div:
        item.req_div, item.req_div_from = parsed.req_div, "폴더 이름"
    elif minutes_files:
        item.req_div, item.req_div_from = "17", "회의록"
        if item.goods:
            item.notes.append("회의록과 물품 표가 함께 있음 — 회의비로 봄(아니면 폴더 이름에 [소모품])")
    elif item.goods:
        item.req_div, item.req_div_from = "85", "물품 표"
    else:
        item.req_div, item.req_div_from = "90", "회의록 · 물품 표 없음"
    if item.req_div == "17" and not item.minutes:
        item.notes.append("회의비인데 회의록이 없음")
    if item.req_div == "85" and not item.goods:
        item.notes.append("소모성물품인데 물품 표(품명 · 수량 · 단가 · 금액 엑셀)가 없음")
    if not item.title:
        item.notes.append("건명이 없음 — 폴더 이름에 사용일 뒤로 적어 주세요")
    return item


def load_job(folder: str, read_text=pdf_text) -> CardJob:
    """결제 기간 폴더 → 비목마다 청구서 · 결제 건. 비목 이름만 있는 하위 폴더는 그 아래를 결제 건으로 본다."""
    folder = os.path.normpath(folder)
    job = CardJob(folder=folder)
    period = parse_period(os.path.basename(folder))
    if period:
        job.start, job.end = period
    items: list[CardItem] = []
    try:
        names = sorted(os.listdir(folder))
    except OSError:
        job.notes.append("폴더를 열지 못함")
        return job
    for name in names:
        path = os.path.join(folder, name)
        if name.startswith((".", "~$")):
            continue
        if not os.path.isdir(path):
            job.skipped.append(f"{name} (기간 폴더 바로 아래 파일 — 결제 건 폴더에 넣어야 붙습니다)")
            continue
        if parse_item_name(name, period).mmdd:
            items.append(read_item(path, period, read_text=read_text))
        elif is_category_folder(name):
            code = _category_of(name)
            subs = sorted(n for n in os.listdir(path) if os.path.isdir(os.path.join(path, n)) and not n.startswith("."))
            for sub in subs:
                if parse_item_name(sub, period).mmdd:
                    items.append(read_item(os.path.join(path, sub), period, code_hint=code, read_text=read_text))
                else:
                    job.skipped.append(f"{name}\\{sub} (이름에 사용일 MMDD 가 없음)")
        else:
            job.skipped.append(f"{name} (이름에 사용일 MMDD 가 없음)")
    by_code: dict[str, CardSheet] = {}
    for item in sorted(items, key=lambda i: (i.used or dt.date.max, i.name)):
        sheet = by_code.setdefault(item.code, CardSheet(item.code, job.start, job.end))
        sheet.items.append(item)
    job.sheets = sorted(by_code.values(), key=lambda s: s.code or "99")
    return job


def looks_like_card_folder(folder: str) -> bool:
    """'자동' 판정 — 이름이 결제 기간이고 안에 결제 건 폴더(사용일 MMDD)가 하나라도 있으면 카드청구 폴더."""
    period = parse_period(os.path.basename(os.path.normpath(folder)))
    if not period:
        return False
    try:
        names = os.listdir(folder)
    except OSError:
        return False
    for name in names:
        path = os.path.join(folder, name)
        if not os.path.isdir(path):
            continue
        if parse_item_name(name, period).mmdd:
            return True
        if is_category_folder(name):
            try:
                if any(parse_item_name(sub, period).mmdd for sub in os.listdir(path)):
                    return True
            except OSError:
                continue
    return False


# ── ERP 미청구 카드내역에서 결제 건 고르기 ─────────────────────────────
@dataclass
class Payment:
    """ERP 카드청구 화면의 미청구 카드내역 한 줄(실측 14절 — k_USED_DATE · k_SHOP_NAME · k_APPRNO · k_USED_COST · k_BAL_AMT · k_CARD_NO)."""
    index: int
    used: dt.date | None = None
    amount: int | None = None      # 사용금액
    balance: int | None = None     # 남은 금액(아직 청구 안 한 것)
    shop: str = ""
    card_last: str = ""
    approval: str = ""
    user: str = ""

    def label(self) -> str:
        day = f"{self.used:%m-%d}" if self.used else "?"
        money = f"{self.amount:,}" if self.amount is not None else "?"
        return f"{day} {self.shop or '?'} {money}원" + (f" (카드 …{self.card_last})" if self.card_last else "")


def _grid_date(text: str) -> dt.date | None:
    digits = re.sub(r"\D", "", str(text or ""))
    return _date(int(digits[:4]), int(digits[4:6]), int(digits[6:8])) if len(digits) >= 8 else None


def payment_from_cells(index: int, cells: dict) -> Payment:
    """그리드 줄의 열 값(k_* → 글자) → Payment. 금액은 title 의 콤마 없는 원본 값일 수도, 콤마 든 글자일 수도 있다."""
    card = re.sub(r"[\s\-]", "", str(cells.get("k_CARD_NO", "") or ""))
    return Payment(index=index, used=_grid_date(cells.get("k_USED_DATE", "")),
                   amount=parse_amount(cells.get("k_USED_COST")), balance=parse_amount(cells.get("k_BAL_AMT")),
                   shop=str(cells.get("k_SHOP_NAME", "") or "").strip(),
                   card_last=card[-4:] if re.search(r"\d{4}$", card) else "",
                   approval=re.sub(r"\D", "", str(cells.get("k_APPRNO", "") or "")),
                   user=str(cells.get("k_USER_NM", "") or "").strip())


def _shop_key(text: str) -> str:
    text = re.sub(r"\(주\)|㈜|주식회사|\(유\)|유한회사", "", str(text or ""))
    return normnm(text).lower()


def shop_match(ours: str, theirs: str) -> bool:
    a, b = _shop_key(ours), _shop_key(theirs)
    return len(a) >= 2 and len(b) >= 2 and (a in b or b in a)


def pick(item: CardItem, payments: list[Payment], taken=()) -> tuple[int | None, str]:
    """결제 건 폴더 하나에 맞는 목록 줄 — 하나로 확실할 때만 그 index, 아니면 (None, 까닭).
    승인번호가 같은 줄 → 아니면 사용일이 같고 남은 금액(또는 사용금액)이 같은 줄 → 여럿이면 카드 뒷자리 · 가맹점으로 좁힌다.
    금액을 모르면 고르지 않는다(그날 결제가 하나뿐이어도 다른 카드 · 다른 건일 수 있다) — 가맹점까지 맞을 때만."""
    pool = [p for p in payments if p.index not in taken and (p.balance is None or p.balance > 0)]
    if item.approval:
        same = [p for p in pool if p.approval and p.approval == item.approval]
        if len(same) == 1 and (item.amount is None or item.amount in (same[0].balance, same[0].amount)):
            return same[0].index, f"승인번호 {item.approval} 가 같음"
    if not item.used:
        return None, "사용일을 몰라 고르지 않음"
    day = [p for p in pool if p.used == item.used]
    if not day:
        return None, f"{item.used:%m-%d} 에 쓴 미청구 결제 건이 목록에 없음"
    cands = day
    if item.amount is not None:
        cands = [p for p in day if item.amount in (p.balance, p.amount)]
        if not cands:
            return None, (f"{item.used:%m-%d} 에 {item.amount:,}원인 결제 건이 없음 — 그날: "
                          + ", ".join(p.label() for p in day[:5]))
    if len(cands) > 1 and item.card_last:
        narrowed = [p for p in cands if p.card_last == item.card_last]
        cands = narrowed or cands
    shop_hit = False
    if item.merchant and (len(cands) > 1 or item.amount is None):
        narrowed = [p for p in cands if shop_match(item.merchant, p.shop)]
        if narrowed:
            cands, shop_hit = narrowed, True
    if len(cands) > 1:
        return None, f"맞는 결제 건이 {len(cands)}개 — " + ", ".join(p.label() for p in cands[:5])
    found = cands[0]
    if item.amount is None and not shop_hit:
        return None, f"금액을 몰라 고르지 않음 — 그날 후보: {found.label()}"
    if item.card_last and found.card_last and found.card_last != item.card_last:
        return None, f"카드 뒷자리가 다름(영수증 …{item.card_last} · 목록 …{found.card_last})"
    why = ["사용일"] + (["금액"] if item.amount is not None else []) + (["가맹점"] if shop_hit else [])
    return found.index, " · ".join(why) + " 이 맞는 줄 하나"


def assign(items: list[CardItem], payments: list[Payment]) -> list[tuple[int | None, str]]:
    """청구서의 결제 건들을 목록 줄에 짝짓는다. 두 폴더가 같은 줄을 고르면 둘 다 사람에게.
    한 번 짝지은 뒤 남은 폴더는 이미 짝지은 줄을 빼고 다시 본다(같은 날 같은 금액 둘 중 하나가 승인번호로 정해진 경우)."""
    result = [pick(item, payments) for item in items]
    for _ in range(2):
        chosen: dict[int, list[int]] = {}
        for i, (index, _why) in enumerate(result):
            if index is not None:
                chosen.setdefault(index, []).append(i)
        for index, who in chosen.items():
            if len(who) > 1:
                for i in who:
                    result[i] = (None, "다른 결제 건 폴더와 같은 줄을 가리킴 — " + ", ".join(items[j].name for j in who))
        taken = {index for index, _w in result if index is not None}
        changed = False
        for i, (index, _why) in enumerate(result):
            if index is None and not result[i][1].startswith("다른 결제 건 폴더"):
                again = pick(items[i], payments, taken)
                if again[0] is not None:
                    result[i] = again
                    taken.add(again[0])
                    changed = True
        if not changed:
            break
    return result
