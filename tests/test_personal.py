"""개인 지급(일반청구 탭 · 증빙 기타) — 지급 공문 읽기 · 명단 모양 · 종류 판정 · 넣기 흐름(담당자 요청 2026-10-19).

실명 · 계좌는 없다. 명단 · 공문 글자는 모양만 같게 코드로 만든다.
"""
from __future__ import annotations

import openpyxl
import pytest

from claim_downloader.claims import kinds, notice, roster, scholarship
from test_scholarship import FakeForm, _patch_popups

NOTICE_TEXT = """수신 내부결재
(경유)
제목 2026년도 9월 재직자 역량 강화 프로그램 지원금 지급
1. 관련: 데이터보안활용혁신융합대학사업단-1234(2026.08.01.)
2. 위 호와 관련하여 재직자 역량 강화 프로그램 지원금을 아래와 같이 지급하고자 합니다.
 가. 지급대상: 재직자 2명
 나. 지급금액: 금540,000원(금오십사만원)
 다. 예산과목: 첨단분야 혁신융합대학사업 / 기업지원·협력활동비
붙임 1. 지원내역 1부. 끝."""


def make_staff(path, rows=None, title="2026년도 9월 재직자 역량 강화 프로그램 지원내역"):
    """재직자 명단 모양(14절): 연번 · 이름 · 사번 · 구분 · 결제자 정보(이름 · 은행 · 계좌 — 가로 묶음) · 금액, 합계 줄."""
    rows = rows or [("가나다", 123456, "재직자", "가나다", "국민", "123-45-678901", 300000),
                    ("라마바", 654321, "재직자", "라마바", "토스", "1000-0000-0001", 240000)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TA장학금"
    ws["B2"] = title
    ws.merge_cells("B2:J2")
    for col, text in zip("BCDE", ("연번", "이름", "사번", "구분")):
        ws[f"{col}3"] = text
        ws.merge_cells(f"{col}3:{col}4")
    ws["F3"] = "결제자 정보(이름·은행·계좌)"
    ws.merge_cells("F3:H3")
    ws["F4"], ws["G4"], ws["H4"] = "이름", "은행", "계좌"
    ws["I3"] = "금액"
    ws.merge_cells("I3:I4")
    r = 5
    for i, row in enumerate(rows, 1):
        for col, value in zip("BCDEFGHI", (i, *row)):
            ws[f"{col}{r}"] = value
        r += 1
    ws[f"B{r}"] = "합계"
    ws[f"I{r}"] = sum(x[-1] for x in rows)
    wb.save(path)
    return path


def make_program(path):
    """프로그램참여지원비 명단 모양(14절): 시트 국외(빈 양식) · 국내, '활동비/여비(… · 합계)' 가로 묶음, 교수 소계 줄, 사번/학번에 탭 문자."""
    wb = openpyxl.Workbook()
    blank = wb.active
    blank.title = "국외"
    ws = wb.create_sheet("국내")
    for sheet in (blank, ws):
        sheet["B2"] = "2026 데이터보안 학술대회 프로그램참여지원비 지원내역"
        for col, text in zip("BCDE", ("연번", "구분", "이름", "사번/학번")):
            sheet[f"{col}3"] = text
            sheet.merge_cells(f"{col}3:{col}4")
        sheet["F3"] = "활동비/여비"
        sheet.merge_cells("F3:J3")
        for col, text in zip("FGHIJ", ("교통비", "일비", "식비", "숙박비", "합계")):
            sheet[f"{col}4"] = text
        sheet["K3"] = "비고"
    rows = [(1, "교수", "홍교수", "123456", 50000, 25000, 25000, 0, 100000), (None, "소계", None, None, None, None, None, None, 100000),
            (2, "학생", "가나다", "\t202600001", 30000, 0, 0, 0, 30000), (3, "학생", "라마바", "202600002", 20000, 0, 0, 0, 20000)]
    for r, row in enumerate(rows, 5):
        for col, value in zip("BCDEFGHIJ", row):
            if value is not None:
                ws[f"{col}{r}"] = value
    ws["B9"] = "합계"
    ws["J9"] = 150000
    wb.save(path)
    return path


def personal_folder(tmp_path, name="재직자 교육 지원(9월)"):
    folder = tmp_path / name
    (folder / "공통증빙").mkdir(parents=True)
    (folder / "개별증빙").mkdir()
    make_staff(folder / "재직자 역량 강화 프로그램 지원내역.xlsx")
    (folder / "공통증빙" / "2026년도 9월 재직자 역량 강화 프로그램 지원금 지급.pdf").write_bytes(b"%PDF-1.4 notice")
    (folder / "공통증빙" / "지원내역.pdf").write_bytes(b"%PDF-1.4 list")
    (folder / "개별증빙" / "가나다 결제내역.pdf").write_bytes(b"%PDF-1.4 a")
    (folder / "개별증빙" / "라마바 결제내역.pdf").write_bytes(b"%PDF-1.4 b")
    return folder


@pytest.fixture
def texts(monkeypatch):
    """PDF 글자 — 파일 이름에 '지급.pdf' 가 든 것만 공문 글자. 나머지는 글자 없음(스캔)."""
    def fake(paths, pages=8):
        return [NOTICE_TEXT if str(p).endswith("지급.pdf") else "" for p in paths if str(p).lower().endswith(".pdf")]
    monkeypatch.setattr(scholarship, "_pdf_texts", fake)
    return fake


# ── 지급 공문 ─────────────────────────────────────────────────────────────
def test_notice_reads_title_category_and_amount():
    paths = ["c/지원내역.pdf", "c/2026년도 9월 재직자 역량 강화 프로그램 지원금 지급.pdf", "c/신청서.pdf"]
    found = notice.find(paths, read_texts=lambda ps: [NOTICE_TEXT if "지급.pdf" in ps[0] else "1. 명단 표"])
    assert found.path.endswith("지급.pdf")
    assert found.title == "2026년도 9월 재직자 역량 강화 프로그램 지원금 지급"
    assert found.clean_title == "2026년도 9월 재직자 역량 강화 프로그램 지원금"       # 끝의 '지급' 은 대표적요에 넣지 않는다
    assert (found.code, found.category) == ("07", "기업지원·협력활동비") and found.amount == 540000


def test_notice_category_forms_and_ambiguity():
    # 보관 폴더식 표기(점 없음) · '그 밖의 사업운영 경비' 도 알아본다
    assert notice.read_category("예산과목: 교육연구프로그램개발운영비")[0] == "03"
    assert notice.read_category("비목 : 그 밖의 사업운영 경비")[0] == "09"
    # 비목 줄이 따로 없으면 글 전체에 하나만 나온 이름
    assert notice.read_category("성과활용·확산지원비로 지급하고자 합니다")[0] == "08"
    # 둘 이상이면 정하지 않는다 — 사람이 고른다
    code, why = notice.read_category("예산과목: 장학금, 교육·연구프로그램개발·운영비")
    assert code == "" and "여럿" in why
    assert notice.read_title("제 목 : 2026 학술대회 등록비 지급(안)\n") == "2026 학술대회 등록비 지급(안)"
    assert notice.Notice(title="2026 학술대회 등록비 지급(안)").clean_title == "2026 학술대회 등록비"
    assert notice.read_title("제목\n2026 학술대회 지원금 지급\n1. 관련") == "2026 학술대회 지원금 지급"


def test_notice_not_found_in_scanned_pdfs():
    found = notice.find(["c/공문.pdf"], read_texts=lambda ps: [""])
    assert not found.path and found.notes


# ── 명단 모양 ─────────────────────────────────────────────────────────────
def test_roster_reads_payer_group_columns(tmp_path):
    r = roster.read(make_staff(tmp_path / "재직자.xlsx"))
    assert [(p.name, p.student_no, p.bank, p.account, p.amount) for p in r.people] == [
        ("가나다", "123456", "국민", "123-45-678901", 300000), ("라마바", "654321", "토스", "1000-0000-0001", 240000)]
    assert r.sheet_total == 540000 and not r.problems()


def test_roster_reads_group_total_column_and_the_filled_sheet(tmp_path):
    r = roster.read(make_program(tmp_path / "프로그램.xlsx"))
    assert [(p.name, p.student_no, p.amount) for p in r.people] == [
        ("홍교수", "123456", 100000), ("가나다", "202600001", 30000), ("라마바", "202600002", 20000)]
    assert r.sheet_total == 150000 and r.total == 150000                     # 소계 줄은 건너뛰었다


def test_roster_without_student_numbers(tmp_path):
    path = tmp_path / "외부.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["외부 참가자 지원내역"])
    ws.append(["연번", "성명", "지원금액", "비고"])
    ws.append([1, "가나다", 50000, ""])
    ws.append(["합계", None, 50000, None])
    wb.save(path)
    r = roster.read(path)
    assert [(p.name, p.amount) for p in r.people] == [("가나다", 50000)]
    assert "학번이 없음" in r.people[0].note           # 입금처는 사람이 고른다


# ── 종류 ──────────────────────────────────────────────────────────────────
def test_general_folder_with_payment_list_becomes_personal(tmp_path):
    folder = personal_folder(tmp_path)
    assert kinds.personal_folder(str(folder))
    assert kinds.guess(str(folder)) == kinds.PERSONAL
    # 담당자: "청구종류 일반으로 해두고 폴더 올리면" — 일반으로 골라도 개인 지급
    assert kinds.resolve(kinds.GENERAL, str(folder)) == kinds.PERSONAL
    # 장학금 이름이 있으면 그대로 장학금, 자문료 · 출장은 아직 일반(증빙 창이 따로 있다)
    sch = tmp_path / "근로장학금" / "9월"
    (sch / "공통증빙").mkdir(parents=True)
    make_staff(sch / "2026학년도 (가)형 근로장학금 지급내역.xlsx")
    assert kinds.guess(str(sch)) == kinds.SCHOLARSHIP and not kinds.personal_folder(str(sch))
    for name in ("자문료 지급(위원회)", "출장비 지급(26.09.01.)"):
        other = tmp_path / name
        (other / "공통증빙").mkdir(parents=True)
        make_staff(other / "지급내역.xlsx")
        assert kinds.guess(str(other)) != kinds.PERSONAL and kinds.resolve(kinds.GENERAL, str(other)) == kinds.GENERAL
    # 묶음 폴더 이름의 '출장' 은 보지 않는다('3. 출장비 및 프로그램참여지원비\…')
    nested = tmp_path / "3. 출장비 및 프로그램참여지원비" / "학술대회"
    (nested / "공통증빙").mkdir(parents=True)
    make_program(nested / "(양식) 프로그램참여지원비 지원내역.xlsx")
    assert kinds.guess(str(nested)) == kinds.PERSONAL
    # 지급내역 엑셀이 공통증빙 안에만 있어도 알아본다
    inside = tmp_path / "WE-Meet 재료비"
    (inside / "공통증빙").mkdir(parents=True)
    make_staff(inside / "공통증빙" / "재료비 지원내역.xlsx")
    assert kinds.guess(str(inside)) == kinds.PERSONAL
    assert scholarship.load_job(str(inside), kind=kinds.PERSONAL).roster.people


def test_personal_job_takes_category_and_summary_from_notice(isolated_home, tmp_path, texts):
    job = scholarship.load_job(str(personal_folder(tmp_path)), kind=kinds.PERSONAL)
    assert job.spec.tab_id == "10" and not job.spec.payroll and job.spec.evidence == "기타"
    assert job.category == "기업지원·협력활동비" and job.pay_category == job.category
    assert job.main_summary == "07_2026년도 9월 재직자 역량 강화 프로그램 지원금"
    assert job.row_summary(0) == "07_2026년도 9월 재직자 역량 강화 프로그램 지원금_가나다"
    assert [len(job.person_files.get(i, [])) for i in range(2)] == [1, 1] and len(job.common_files) == 2
    assert not job.problems(), job.problems()                          # 지급예정년월은 묻지 않는다(인건비 탭 칸)
    lines = "\n".join(job.summary_lines())
    assert "지급 공문" in lines and "증빙구분 기타" in lines and "지급예정년월" not in lines
    # 비목을 바꾸면 대표적요 앞 번호도
    job.set_category("교육·연구프로그램개발·운영비")
    assert job.main_summary.startswith("03_2026년도 9월")
    # 공문 금액과 명단 합계가 다르면 알린다(막지는 않는다)
    job.notice.amount = 500000
    assert any("지급 공문의 금액" in p for p in job.problems())


def test_personal_job_without_notice_asks_for_category(isolated_home, tmp_path, monkeypatch):
    monkeypatch.setattr(scholarship, "_pdf_texts", lambda paths, pages=8: ["" for _ in paths])
    job = scholarship.load_job(str(personal_folder(tmp_path)), kind=kinds.PERSONAL)
    assert job.category == "" and job.main_summary.startswith("00_")
    problems = job.problems()
    assert any(p.startswith("비목이 없습니다") for p in problems)
    assert any("지급 공문" in p for p in problems)
    job.set_category("기업지원·협력활동비")
    assert job.main_summary.startswith("07_") and not any(p.startswith("비목") for p in job.problems())


# ── 넣기 ──────────────────────────────────────────────────────────────────
def test_personal_run_fills_general_tab_with_other_evidence(isolated_home, tmp_path, texts, monkeypatch):
    job = scholarship.load_job(str(personal_folder(tmp_path)), kind=kinds.PERSONAL)
    job.tab = scholarship.person.TAB_STAFF
    frame = FakeForm(account="", owner_result="")      # 인적정보에 계좌가 없는 사람 — 명단(결제자 정보)의 계좌를 넣는다
    calls = _patch_popups(monkeypatch, frame)
    logs = []
    report = scholarship.run(object(), frame, job, logs.append, ask=lambda m: False, wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and not report.stopped, report.lines()
    assert scholarship.TAB_ID not in frame.clicks and scholarship.GENERAL_TAB_ID not in frame.clicks   # 이미 일반청구 탭
    assert [c[2] for c in calls["person"]] == [scholarship.person.TAB_STAFF] * 2
    first = frame.saved[0]
    assert first["PAY_EXP_CD"] == "기업지원·협력활동비" and first["EVDI_DIV_CD"] == "기타"
    assert "INCM_DIV_CD" not in first and "PAY_PLAN_YM" not in first     # 인건비 탭 칸은 건드리지 않는다
    assert first["REQ_AMT"] == "300000" and first["SUPPLY_AMT"] == "300000" and first["TAX_AMT"] == "0"
    assert first["REQ_PTCL"] == job.main_summary + "_가나다" and first["APPL_CONT"] == job.main_summary
    assert first["ACCT_NO"] == "12345678901"                            # 결제자 정보의 계좌
    assert frame.clicks.count(scholarship.OWNER_CHECK_ID) == 2          # 계좌를 넣었으니 예금주조회를 새로
    assert not any("국적" in line for line in logs)


def test_personal_run_stops_when_claim_type_is_not_other(isolated_home, tmp_path, texts, monkeypatch):
    job = scholarship.load_job(str(personal_folder(tmp_path)), kind=kinds.PERSONAL)
    frame = FakeForm()
    frame.v["REQ_DIV"] = "33"                                           # 여비(국내)로 바뀌어 있다
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 0 and report.stopped and "(90) 기타" in asked[0]
    assert not any(c.startswith("저장:") for c in frame.clicks)


def test_personal_run_refuses_without_category(isolated_home, tmp_path, monkeypatch):
    monkeypatch.setattr(scholarship, "_pdf_texts", lambda paths, pages=8: ["" for _ in paths])
    job = scholarship.load_job(str(personal_folder(tmp_path)), kind=kinds.PERSONAL)
    frame = FakeForm()
    _patch_popups(monkeypatch, frame)
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: False, wait=lambda s: None, dry_run=False)
    assert report.saved == 0 and "비목" in report.results[0].note


def test_personal_plan_marks_test_mode(isolated_home, tmp_path, texts):
    job = scholarship.load_job(str(personal_folder(tmp_path)), kind=kinds.PERSONAL)
    job.test_mode = True
    with scholarship.marked(job):
        report = scholarship.run(None, None, job, lambda m: None, ask=lambda m: False, dry_run=True)
    assert all(r.summary.startswith("[시험] 07_") for r in report.results)


def test_notice_title_is_summary_only_when_it_is_a_payment_notice(isolated_home, tmp_path, monkeypatch):
    """담당자 실제(2026-10-20): 10월 폴더의 공문이 '… 지도교수 추가'(계획 변경)였는데 대표적요가 그 제목이 됐다 — 담당자가 명단 제목으로 고쳤다.
    '지급' 공문이 아니면 명단 제목으로 만든다(비목은 그 공문에서 그대로)."""
    plan = NOTICE_TEXT.replace("제목 2026년도 9월 재직자 역량 강화 프로그램 지원금 지급", "제목 2026학년도 재직자 역량 강화 [AI 바이브 코딩] 지도교수 추가")
    monkeypatch.setattr(scholarship, "_pdf_texts",
                        lambda paths, pages=8: [plan if str(p).endswith("지급.pdf") else "" for p in paths if str(p).lower().endswith(".pdf")])
    job = scholarship.load_job(str(personal_folder(tmp_path)), kind=kinds.PERSONAL)
    assert job.category == "기업지원·협력활동비"
    assert job.main_summary == "07_2026년도 9월 재직자 역량 강화 프로그램 지원내역"


def test_reload_keeps_edited_summary_category_and_unchecked_people(isolated_home, tmp_path, texts):
    """다시 읽기([다시 읽기] · 엑셀 · 종류)에도 고친 대표적요 · 비목, 넣어서 체크를 푼 사람은 남는다(2026-10-20 · 10-21 실제)."""
    folder = str(personal_folder(tmp_path))
    old = scholarship.load_job(folder, kind=kinds.PERSONAL)
    old.main_summary = "07_재직자 역량 강화 프로그램_AI 바이브 코딩 교육 지원내역(10월)"
    old.selected[0] = False
    new = scholarship.load_job(folder, kind=kinds.PERSONAL)
    scholarship.keep_edits(new, old)
    assert new.main_summary == old.main_summary and new.selected == [False, True]
    # 고치지 않았으면 새로 읽은 값(공문 · 명단이 바뀌었을 수 있다)
    fresh = scholarship.load_job(folder, kind=kinds.PERSONAL)
    again = scholarship.load_job(folder, kind=kinds.PERSONAL)
    scholarship.keep_edits(again, fresh)
    assert again.main_summary == fresh.default_summary and again.selected == [True, True]
    # 비목을 바꾸면 그 비목과 바뀐 대표적요 번호가 남는다
    fresh.set_category("교육·연구프로그램개발·운영비")
    again = scholarship.load_job(folder, kind=kinds.PERSONAL)
    scholarship.keep_edits(again, fresh)
    assert again.category == "교육·연구프로그램개발·운영비" and again.main_summary.startswith("03_")
