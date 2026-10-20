"""카드청구 넣기 준비물 — 회의록 표(참석자 · 시각 · 내용) · 참여연구원 명단 · 참석자 나누기 · 계획 · 이 청구서인지 · 기록.

담당자 규칙(2026-10-20): 회의록의 소속이 비었거나 · 강원대학교 · 학과명이면 참여인력이나 내부참석자, 그 밖의 소속은 외부참석자.
회의비 · 물품 줄 적요 앞 `NN_` 은 건별 공문의 비목. 실명 · 직번은 없다 — 모양만 같은 가짜 값이다.
ERP 화면을 실제로 누르는 흐름은 모의 카드청구 화면 + Chromium 으로 확인했다(CLAUDE.md 19절).
"""
from __future__ import annotations

import datetime as dt

import openpyxl

from claim_downloader.claims import card, card_erp, participants
from claim_downloader.claims.form import STATE_JS
from claim_downloader.minutes.docx import build_docx
from claim_downloader.minutes.hwpx import build_hwpx

from test_card import RECEIPT, make_goods, minutes_data, reader, touch

D = dt.date

PDF_TEXT = """참여연구원 현황표
과제번호 202400000000 사업명 가짜사업
연구책임자 성명 김책임 소속 가나다학과
당해사업비 1,000,000 당해연구기간 2026-03-01 ~ 2027-02-28
성명 소속 역할 직급 학위과정 소득구분 지급비목 지급구분
참여시작일 참여종료일 개월 계상률 월지급액 총지급액 지급방법
김책임 가나다학과 책임연구원 교수 박사과정졸업
(100001) 0.00 0.00 % 0 0
이연구 산학협력단_근로
장학생 공동연구원 연구원 석사과정졸업 근로소득 인건비 현금
(900001) 2026-08-01 2027-02-28 7.00 100.00 % 1,000 7,000 계좌입금
박학생
데이터보안·활용혁
신융합대학사업단 보조연구원 연구원 대학(교)졸업
1/4
(202600001) 2026-04-01 2027-02-28 11.00 100.00 % 1,000 11,000 계좌입금
최긴번 가나다학 보조연구원 대학원생(석사) 석사과정재학
(2024000000001)
0.00 0.00 % 0 0
합계(미지급,현물 제외) 18,000
2/4"""


# ── 참여연구원 명단 ────────────────────────────────────────────────
def test_participant_pdf_two_lines_per_person_even_across_pages():
    """한 사람이 두 줄(윗줄 이름 · 소속 …, 아랫줄 (직번) …). 소속이 꺾이거나 쪽이 바뀌어도, 직번이 6 · 9 · 13자리여도."""
    got = participants.read_pdf_text(PDF_TEXT)
    assert [(m.name, m.emp_no) for m in got] == [("김책임", "100001"), ("이연구", "900001"), ("박학생", "202600001"),
                                                 ("최긴번", "2024000000001")]
    assert got[0].dept == "가나다학과" and got[1].dept == "산학협력단_근로"


def test_participant_xlsx_name_then_emp_no_row(tmp_path):
    """ERP 출력 엑셀(Prn_PrjPartInfo.xlsx, 14절): A열 윗줄 이름 · B열 소속, 아랫줄 A열 '(직번)'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A9"], ws["B9"] = "성명\n(직번)", "소속"
    ws["A11"], ws["B11"], ws["E11"] = "김책임", "가나다학과", "책임연구원"
    ws["A12"], ws["B12"] = "(100001)", None
    ws["A13"], ws["B13"] = "이연구", "산학협력단_근로장학생"
    ws["A14"], ws["B14"] = "(900001)", "2026-08-01"
    path = tmp_path / "Prn_PrjPartInfo.xlsx"
    wb.save(path)
    got = participants.read(str(path))
    assert [(m.name, m.emp_no, m.dept) for m in got] == [("김책임", "100001", "가나다학과"), ("이연구", "900001", "산학협력단_근로장학생")]
    assert participants.find_file("", str(tmp_path)) == str(path)              # 설정이 비면 폴더에서 찾는다
    assert participants.find_file(str(path)) == str(path)


def test_attendee_rule_blank_knu_or_department_is_inside():
    """담당자 규칙: 소속이 공란 · 강원대학교 · 학과명이면 참여인력/내부참석자, 그 밖은 외부참석자."""
    for org in ("", "강원대학교", "강원대 AI융합학과", "컴퓨터공학과", "정보보호학", "데이터보안·활용혁신융합대학사업단", "산학협력단", "공과대학"):
        assert participants.is_internal_org(org), org
    for org in ("한림대학교", "(주)가나다", "가나다 주식회사", "한국가나다재단", "춘천시청", "한국전자통신연구원"):
        assert not participants.is_internal_org(org), org
    members = [participants.Member("김책임", "100001"), participants.Member("이연구", "900001"), participants.Member("이연구", "900002")]
    got = participants.classify([("강원대학교", "김책임"), ("", "홍길동"), ("한림대학교", "김책임"), ("가나다학과", "이연구")], members)
    assert [(a.kind, a.emp_no) for a in got] == [("part", "100001"), ("internal", ""), ("external", ""), ("internal", "")]
    assert "외부참석자로" in got[2].note and "2명" in got[3].note


# ── 회의록 표 ────────────────────────────────────────────────────────
def test_minutes_table_reads_merged_orgs_times_and_content(tmp_path):
    """소속 칸은 같은 소속끼리 세로로 합쳐져(hwpx rowSpan · docx vMerge) 글자 순서로는 못 가린다 — 칸 주소로 읽는다."""
    data = minutes_data(attendees=[{"org": "강원대학교", "name": "가"}, {"org": "강원대학교", "name": "나"}, {"org": "한림대학교", "name": "다"}],
                        attendeeCount=5, date="2026. 9. 15. (화) 9:30 ~ 11:00",
                        content=["첫째 문단을 논의하였음.", "둘째 문단을 정하였음."])
    for build, name in ((build_hwpx, "m.hwpx"), (build_docx, "m.docx")):
        path = tmp_path / name
        build(data, str(path))
        got = card.read_minutes(str(path))
        assert got.attendees == [("강원대학교", "가"), ("강원대학교", "나"), ("한림대학교", "다")], name   # 빈 칸 둘은 뺀다
        assert (got.start, got.end) == ("09:30", "11:00"), name
        assert got.agenda == "사업단 운영회의" and got.place == "사업단 회의실" and got.used == D(2026, 9, 15)
        assert got.content == ["첫째 문단을 논의하였음.", "둘째 문단을 정하였음."], name


# ── 계획 · 적요 ──────────────────────────────────────────────────────
def _meeting_item(tmp_path, notice_text="제목 운영회의 회의비\n예산과목: 기타운영경비\n장소: 사업단 회의실(춘천)", **over):
    period = tmp_path / "26-09(09.09. ~ 10.08.)"
    folder = period / "기타운영경비 0915 사업단 운영회의"
    touch(folder / "0915 영수증.pdf")
    touch(folder / "회의비 지출 공문.pdf")
    touch(folder / "사진.jpg")
    build_hwpx(minutes_data(**over), str(folder / "2026-09-15 화 운영회의.hwpx"))
    job = card.load_job(str(period), read_text=reader({"0915 영수증.pdf": RECEIPT, "회의비 지출 공문.pdf": notice_text}))
    return job, job.items[0]


def test_meeting_line_prefix_comes_from_the_notice(tmp_path):
    """담당자(2026-10-20): 회의비 줄 적요는 ERP 가 회의록에서 채운 글 앞에 '건별로 공문을 확인하고 공문에 있는 비목' 을 붙인다."""
    job, item = _meeting_item(tmp_path)
    assert item.req_div == "17" and item.notice_code == "09" and item.line_code == "09"
    assert item.prefixed("회의일자:2026-09-15/ 회의장소:사업단 회의실") == "09_회의일자:2026-09-15/ 회의장소:사업단 회의실"
    assert item.prefixed("09_회의일자:…") == "09_회의일자:…" and item.prefixed("") == ""
    job, item = _meeting_item(tmp_path / "b", notice_text="제목 운영회의\n예산과목: 기업지원·협력활동비")
    assert item.code == "09" and item.line_code == "07" and any("공문의 비목(07)" in n for n in item.notes)


def test_plan_shows_who_the_app_types_and_who_the_person_adds(tmp_path):
    job, item = _meeting_item(tmp_path, attendees=[{"org": "강원대학교", "name": "가"}, {"org": "가나다학과", "name": "나"},
                                                    {"org": "한림대학교", "name": "다"}],
                              content=["운영 현황과 일정을 함께 논의하였음.", "다음 회의 날짜를 정하였음."])
    plan = card_erp.plan_item(item, 0, [participants.Member("가", "100001")])
    text = "\n".join(plan.lines)
    assert "회의록등록: 11:00~12:00 · 장소 사업단 회의실" in text
    assert "참여인력: 가(100001)" in text and "내부참석자: 나 · 가나다학과 — 사람이 돋보기로" in text
    assert "외부참석자: 다 · 한림대학교 — 앱이 넣음" in text and "09_" in text
    assert card_erp.attachments(item)[0].endswith("회의비 지출 공문.pdf")          # 공문이 맨 앞, PDF · 그림만
    assert {p.rsplit("/", 1)[-1] for p in card_erp.attachments(item)} == {"회의비 지출 공문.pdf", "0915 영수증.pdf", "사진.jpg"}
    _v, _p, missing = card_erp.meeting_plan(item, [])
    assert not missing
    job, short = _meeting_item(tmp_path / "s", content=["짧음."], attendees=[])
    _v, _p, missing = card_erp.meeting_plan(short, [])
    assert "회의내용이 20자보다 짧음" in missing and "참석자 이름" in missing


def test_plan_for_other_and_goods(tmp_path):
    period = tmp_path / "26-09"
    touch(period / "기타운영경비 0920 소프트웨어 구독 33,000원" / "결제.pdf")
    (period / "기타운영경비 0925 사무용품").mkdir(parents=True)
    make_goods(period / "기타운영경비 0925 사무용품" / "구매의뢰.xlsx")
    job = card.load_job(str(period), read_text=reader({}))
    other, goods = job.items
    assert "청구내역: 09_소프트웨어 구독" in card_erp.plan_item(other, 0).lines
    assert any("사람이 넣음" in line for line in card_erp.plan_item(goods, 1).lines)
    assert card_erp.main_summary(job.sheets[0], test_mode=True).startswith("[시험] 09_[그밖의사업운영경비] 카드 결제 건")


# ── 이 청구서인지 · 기록 ───────────────────────────────────────────────
class FakeFrame:
    def __init__(self, **values):
        self.values = values

    def evaluate(self, js, arg=None):
        assert js == STATE_JS
        return {k: {"tag": "input", "value": self.values.get(k, "")} for k in arg}


def test_sheet_owner_new_same_number_or_same_summary():
    sheet = card.CardSheet("09", D(2026, 9, 9), D(2026, 10, 8))
    summary = sheet.main_summary
    assert card_erp.sheet_owner(FakeFrame(), sheet, summary)[0]                                    # 새 카드청구서
    ok, why = card_erp.sheet_owner(FakeFrame(REQ_CNT="335", APPR_DIV_CD="40"), sheet, summary)
    assert not ok and "335" in why                                                                  # 남의 청구서
    sheet.req_cnt = "335"
    assert card_erp.sheet_owner(FakeFrame(REQ_CNT="335", APPR_DIV_CD="40"), sheet, summary)[0]     # 이 청구서(결의번호)
    sheet.req_cnt = ""
    assert card_erp.sheet_owner(FakeFrame(REQ_CNT="336", APPR_DIV_CD="40", APPL_CONT=summary), sheet, summary)[0]


def test_record_and_restore_saved_items(tmp_path):
    period = tmp_path / "26-09"
    touch(period / "기타운영경비 0920 소프트웨어 구독 33,000원" / "결제.pdf")
    touch(period / "기타운영경비 0921 도메인 연장 22,000원" / "결제.pdf")
    job = card.load_job(str(period), read_text=reader({}))
    sheet = job.sheets[0]
    report = card_erp.CardReport(dry_run=False, summary=sheet.main_summary, req_cnt="335",
                                 results=[card_erp.ItemResult(0, sheet.items[0].name, "saved")])
    sheet.items[0].status = "saved"
    card_erp.record_sheet(job, sheet, report)
    again = card.load_job(str(period), read_text=reader({}))
    card_erp.restore(again)
    assert again.sheets[0].req_cnt == "335"
    assert [i.status for i in again.sheets[0].items] == ["saved", ""]


def test_dry_run_plans_every_item_without_a_browser(tmp_path):
    job, item = _meeting_item(tmp_path)
    report = card_erp.run(None, None, job, job.sheets[0], lambda _m: None, ask=lambda _m: False, dry_run=True)
    assert report.dry_run and [r.status for r in report.results] == ["planned"] and not report.blocked
    assert report.lines()[0] == "미리 보기 (ERP 에 넣지 않음)"
