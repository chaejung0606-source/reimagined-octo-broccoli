"""장학금 청구(claims/roster.py · person.py · attach.py · scholarship.py).

실제 명단·청구서 없이 돈다 — 명단 엑셀은 openpyxl 로 가짜 이름·학번으로 만든다(10절). ERP 는 evaluate 가 JS 상수를
알아보고 화면처럼 반응하는 가짜 프레임으로 흉내 낸다. 실제 화면 모양은 CLAUDE.md 14절 '장학금' ①~⑤.
"""
from __future__ import annotations

import os
import pathlib

import pytest

openpyxl = pytest.importorskip("openpyxl")

from claim_downloader.claims import attach, form, nation, person, roster, scholarship  # noqa: E402
from claim_downloader.claims import payee  # noqa: E402


# ── 가짜 명단 ─────────────────────────────────────────────────────────
def make_ga(path, rows=None, total=None, title="2026년도 7월 (가)형 근로장학금 지급내역"):
    """(가)형 모양 — 실측(2026-10-14): 시트 'TA장학금', B2 제목, 3~4행 머리글(산출내역 병합), 5행부터 학생, 합계 행."""
    rows = rows or [("가나다", 202600001, "국민", "123-45-678901", 12000, "25", 300000),
                    ("라마바", 202600002, "토스", "1000-0000-0001", 12000, "20", 240000)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TA장학금"
    ws["B2"] = title
    ws.merge_cells("B2:I2")
    for col, text in zip("BCDEF", ("연번", "이   름", "학   번", "은행", "계   좌")):
        ws[f"{col}3"] = text
        ws.merge_cells(f"{col}3:{col}4")
    ws["G3"] = "산출내역"
    ws.merge_cells("G3:I3")
    ws["G4"], ws["H4"], ws["I4"] = "단가(A)", "시간(B)", "금액(C=A*B)"
    r = 5
    for i, (name, no, bank, acct, unit, hours, amount) in enumerate(rows, 1):
        for col, value in zip("BCDEFGHI", (i, name, no, bank, acct, unit, hours, amount)):
            ws[f"{col}{r}"] = value
        r += 1
    ws[f"B{r}"] = "합   계"
    ws.merge_cells(f"B{r}:H{r}")
    ws[f"I{r}"] = total if total is not None else sum(x[-1] for x in rows)
    wb.save(path)
    return path


def make_na(path):
    """(나)형 TA 모양 — 강좌명 열이 끼고, 한 학생이 두 강좌면 두 줄 + 이름·학번·은행·계좌 세로 병합."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TA장학금"
    ws["B2"] = "2026년도 6월 (나)형 근로장학금 지급내역"
    for col, text in zip("BCDEFG", ("연번", "이름", "학번", "강좌명", "은행", "계좌")):
        ws[f"{col}3"] = text
        ws.merge_cells(f"{col}3:{col}4")
    ws["H3"] = "산출내역"
    ws.merge_cells("H3:J3")
    ws["H4"], ws["I4"], ws["J4"] = "단가(A)", "시간(B)", "금액(C=A*B)"
    ws["M3"], ws["N3"] = "연락처", "이메일"
    data = [(1, "사아자", 202600010, "자료구조", "국민은행", "111-11-111111", 15000, "10", 150000),
            (2, None, None, "운영체제", None, None, 15000, "8", 120000),
            (3, "차카타", 202600011, "알고리즘", "토스뱅크", "1000-1111-2222", 20000, "5", 100000)]
    for i, row in enumerate(data):
        for col, value in zip("BCDEFGHIJ", row):
            ws[f"{col}{5 + i}"] = value
    ws.merge_cells("C5:C6")
    ws.merge_cells("D5:D6")
    ws["B8"] = "합계"
    ws["J8"] = 370000
    wb.save(path)
    return path


def add_nation(path, values, header="국   적", col="K"):
    """명단에 국적 열을 더한다 — 담당자가 외국인 학생 칸만 채우기로 함(2026-10-14). values: 엑셀 행 → 국적."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws[f"{col}3"] = header
    ws.merge_cells(f"{col}3:{col}4")
    for row, value in values.items():
        ws[f"{col}{row}"] = value
    wb.save(path)
    return path


def test_roster_reads_nation_column_and_fills_merged_rows(tmp_path):
    r = roster.read(add_nation(make_ga(tmp_path / "가.xlsx"), {6: "우즈베키스탄"}))
    assert [p.nation for p in r.people] == ["", "우즈베키스탄"]
    assert [p.foreign_nation for p in r.people] == ["", "우즈베키스탄"] and not r.problems()
    path = make_na(tmp_path / "나.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws["K3"], ws["K5"], ws["K7"] = "국적", "베트남", "대한민국"
    ws.merge_cells("K5:K6")                                   # 두 강좌 학생 — 국적 칸도 세로 병합
    wb.save(path)
    r = roster.read(path)
    assert [p.nation for p in r.people] == ["베트남", "베트남", "대한민국"]
    assert [p.foreign_nation for p in r.people] == ["베트남", "베트남", ""]       # 한국은 외국인이 아니다


def test_nation_header_is_not_confused_with_other_columns():
    assert roster._header_key("국적") == "nation" and roster._header_key("국적(국가/지역)") == "nation"
    assert roster._header_key("국가장학금") is None


def _nations(*pairs):
    return [{"index": i, "text": f"{code} {name}", "cells": {"k_NATION_CD": code, "k_NATION_NM": name}}
            for i, (code, name) in enumerate(pairs)]


def test_nation_match_picks_only_one_sure_row():
    """실측 18:46 — '우즈' 검색에 UZ 우즈베키스탄 한 줄. 같은 이름이 하나면 고르고, 아니면 사람."""
    assert nation.match(_nations(("UZ", "우즈베키스탄")), "우즈베키스탄") == (0, "국적 이름 일치")
    assert nation.match(_nations(("UZ", "우즈베키스탄")), "우즈베키스탄 공화국")[0] == 0
    assert nation.match(_nations(("US", "미국"), ("AS", "미국령사모아")), "미국") == (0, "국적 이름 일치")
    assert nation.match(_nations(("VN", "베트남사회주의")), "베트남")[0] == 0          # 이름을 담은 줄이 하나
    assert nation.match(_nations(("CN", "중국")), "cn") == (0, "국적 이름 일치")        # 두 글자 코드로 적어도
    index, why = nation.match(_nations(("CD", "콩고민주공화국"), ("CG", "콩고공화국")), "콩고")
    assert index is None and "2줄" in why
    index, why = nation.match(_nations(("GH", "가나")), "몽골")
    assert index is None and "없음" in why
    assert nation.match([], "몽골") == (None, "검색 결과가 없음")


def test_nation_queries_try_short_name_too():
    assert nation.queries_for("우즈베키스탄") == ["우즈베키스탄"]
    assert nation.queries_for("우즈베키스탄 공화국") == ["우즈베키스탄 공화국", "우즈베키스탄"]


def test_same_country_ignores_spaces_and_republic():
    assert scholarship.same_country("우즈베키스탄", "우즈베키스탄 공화국")
    assert scholarship.same_country("중국", "중화인민공화국") is False          # 다르게 적으면 사람이 본다(신청 전 확인)
    assert not scholarship.same_country("베트남", "") and not scholarship.same_country("베트남", "몽골")


def test_roster_reads_ga_layout_by_header_words(tmp_path):
    r = roster.read(make_ga(tmp_path / "명단.xlsx"))
    assert r.title == "2026년도 7월 (가)형 근로장학금 지급내역" and r.kind == "가"
    assert [p.name for p in r.people] == ["가나다", "라마바"]
    assert [p.student_no for p in r.people] == ["202600001", "202600002"]
    assert [p.amount for p in r.people] == [300000, 240000]       # '금액(C=A*B)' 열 — 단가·시간 열이 아니다
    assert r.people[0].account_digits == "12345678901" and r.people[1].bank == "토스"
    assert r.sheet_total == 540000 and r.total == 540000 and r.problems() == []


def test_roster_fills_merged_second_row_and_reads_course(tmp_path):
    r = roster.read(make_na(tmp_path / "TA.xlsx"))
    assert r.kind == "나"
    assert [(p.name, p.student_no, p.course, p.amount) for p in r.people] == [
        ("사아자", "202600010", "자료구조", 150000), ("사아자", "202600010", "운영체제", 120000),
        ("차카타", "202600011", "알고리즘", 100000)]
    assert r.people[1].account == "111-11-111111"                  # 병합된 둘째 줄은 앞 줄 값을 내려받는다
    assert r.sheet_total == 370000 and not any("합계" in p for p in r.problems())
    # 학번까지 같고 교과목이 다르면 한 학생이 교과목 둘을 맡은 것 — 문제가 아니다(담당자 규칙 2026-10-14)
    assert r.duplicate_problems() == [] and not any("중복" in p for p in r.problems())


def _people(*rows):
    return [roster.Person(name=n, student_no=no, course=c, amount=1, row=5 + i) for i, (n, no, c) in enumerate(rows)]


def test_same_name_is_told_apart_by_student_no_then_course():
    ok = roster.Roster(path="", people=_people(("가나다", "202600001", "자료구조"),
                                                ("가나다", "202600002", "자료구조"),     # 동명이인 — 학번이 다르다
                                                ("가나다", "202600001", "운영체제")))    # 같은 학생의 다른 교과목
    assert ok.duplicate_problems() == []
    twice = roster.Roster(path="", people=_people(("가나다", "202600001", "자료구조"), ("가나다", "202600001", "자료구조")))
    assert any("중복" in p and "자료구조" in p for p in twice.duplicate_problems())
    no_course = roster.Roster(path="", people=_people(("가나다", "202600001", ""), ("가나다", "202600001", "")))
    assert any("중복" in p for p in no_course.duplicate_problems())
    wrong = roster.Roster(path="", people=_people(("가나다", "202600001", ""), ("라마바", "202600001", "")))
    assert any("이름이 다릅니다" in p for p in wrong.duplicate_problems())


def test_roster_flags_total_mismatch_and_missing_student_no(tmp_path):
    rows = [("가나다", None, "국민", "1", 10000, "1", 10000), ("라마바", 202600002, "국민", "2", 10000, "2", 20000)]
    r = roster.read(make_ga(tmp_path / "x.xlsx", rows=rows, total=99999))
    problems = " / ".join(r.problems())
    assert "합계" in problems and "학번이 없음" in problems


def test_summaries_follow_type_rules():
    rules = scholarship.parse_rules(scholarship.RULES_DEFAULT)
    main = scholarship.main_summary("2026년도 7월 (가)형 근로장학금 지급내역", rules)
    assert main == "02_2026년도 7월 (가)형 근로장학금"             # '지급내역' 은 뺀다 (담당자 지정 2026-10-14)
    assert scholarship.main_summary("2026년도 7월 (가)형 근로장학금 지급 내역", rules) == main
    ga = roster.Person(name="가나다", student_no="202600001", amount=1)
    assert scholarship.row_summary(main, ga, "가", rules) == main + "_가나다"
    # TA 근로장학금: 대표적요(이름_교과목명) — 담당자 지정 2026-10-14. 제목에 '(나)형' 이 없어도 교과목이 있으면 이 꼴
    na = roster.Person(name="사아자", student_no="202600010", course="자료구조", amount=1)
    ta_main = scholarship.main_summary("2026년도 9월 (나)형 근로장학금 지급내역", rules)
    assert scholarship.row_summary(ta_main, na, "나", rules) == "02_2026년도 9월 (나)형 근로장학금(사아자_자료구조)"
    assert scholarship.row_summary("02_X", na, "", rules) == "02_X(사아자_자료구조)"
    # 유형 글자 규칙을 담당자가 따로 적었으면 그것이 먼저
    typed = scholarship.parse_rules(scholarship.RULES_DEFAULT + "\n행적요(나): {대표적요}({이름}_{학번}_{강좌명})")
    assert scholarship.row_summary("02_X", na, "나", typed) == "02_X(사아자_202600010_자료구조)"
    custom = scholarship.parse_rules("대표적요: 02_[선지원금] {제목}\n행적요: {대표적요}-{이름}")
    assert scholarship.main_summary("T", custom) == "02_[선지원금] T"
    assert scholarship.main_summary("T 지급내역", custom) == "02_[선지원금] T 지급내역"   # '뺄 말' 이 없는 규칙이면 제목 그대로


def test_load_job_splits_common_and_person_files(tmp_path):
    folder = tmp_path / "7월 (가)형"
    folder.mkdir()
    make_ga(folder / "2026학년도 (가)형 근로장학금 지급내역.xlsx")
    common = folder / "공통증빙"
    common.mkdir()
    (common / "지급 공문.pdf").write_bytes(b"%PDF")
    (common / "명단.pdf").write_bytes(b"%PDF")
    each = folder / "개별증빙"
    each.mkdir()
    (each / "가나다.pdf").write_bytes(b"%PDF")
    (each / "라마바").mkdir()
    (each / "라마바" / "근무일지.pdf").write_bytes(b"%PDF")
    (each / "라마바" / "청구서.pdf").write_bytes(b"%PDF")
    (each / "근무일지 등.zip").write_bytes(b"PK")                  # 압축은 올리지 않는다
    (each / "누군지모름.pdf").write_bytes(b"%PDF")
    job = scholarship.load_job(str(folder))
    assert job.roster.kind == "가" and len(job.people) == 2 and job.selected == [True, True]
    assert sorted(f.split("/")[-1].split("\\")[-1] for f in job.common_files) == ["명단.pdf", "지급 공문.pdf"]
    assert [len(job.person_files.get(i, [])) for i in range(2)] == [1, 2]
    assert len(job.files_for(1)) == 4                              # 공통 2 + 개별 2
    assert [f.rsplit("/", 1)[-1] for f in job.unmatched] == ["누군지모름.pdf"]
    assert any("모르는 개별증빙" in p for p in job.problems())
    assert job.main_summary.startswith("02_2026년도 7월 (가)형")


def test_old_default_rules_file_is_upgraded_but_edited_one_is_kept():
    from claim_downloader import settings
    old = scholarship.RULES_HEADER + "대표적요: 02_{제목}\n행적요: {대표적요}_{이름}\n행적요(나): {대표적요}({이름}_{학번}_{강좌명})\n"
    path, base = settings.scholarship_rules_path(), settings.scholarship_rules_base_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(old, encoding="utf-8")
    base.write_text(old, encoding="utf-8")
    rules = scholarship.load_rules()                      # 손대지 않은 옛 기본값 → 새 기본값으로
    assert "행적요(강좌)" in rules and "행적요(나)" not in rules
    assert rules["행적요(강좌)"] == "{대표적요}({이름}_{강좌명})" and "지급내역" in rules["제목에서뺄말"]
    path.write_text(old.replace("02_{제목}", "02_[선지원금] {제목}"), encoding="utf-8")
    rules = scholarship.load_rules()                      # 담당자가 고친 것 → 그대로
    assert rules["대표적요"] == "02_[선지원금] {제목}" and "행적요(나)" in rules


def test_job_warns_when_two_rows_get_the_same_summary(tmp_path):
    rows = [("가나다", 202600001, "국민", "1", 10000, "1", 10000), ("가나다", 202600002, "국민", "2", 10000, "1", 10000)]
    folder = tmp_path / "7월"
    folder.mkdir()
    make_ga(folder / "명단.xlsx", rows=rows, total=20000)
    job = scholarship.load_job(str(folder))
    assert any("개별 적요가 같은 줄" in p for p in job.problems())      # (가)형 동명이인은 '대표적요_이름' 이 같아진다


def test_ta_student_with_two_courses_gets_files_per_course_or_both(tmp_path):
    people = [roster.Person("사아자", "202600010", course="자료구조"), roster.Person("사아자", "202600010", course="운영체제"),
              roster.Person("차카타", "202600011", course="알고리즘")]
    files = [str(tmp_path / "사아자.pdf"), str(tmp_path / "사아자_운영체제 근무일지.pdf"), str(tmp_path / "차카타.pdf")]
    assigned, unmatched = scholarship.assign_person_files(files, str(tmp_path), people)
    # 교과목이 없는 개인 서류는 그 학생의 두 줄 모두에, 교과목이 적힌 것은 그 줄에만
    assert assigned == {0: [files[0]], 1: [files[0], files[1]], 2: [files[2]]} and unmatched == []
    job = scholarship.Job(folder=str(tmp_path), roster=roster.Roster(path="", people=people), selected=[True, True, True])
    assert job.count_text() == "3줄(2명)" and job.count_text([0, 2]) == "2명"
    # 결과 표(ui/report_dialog)가 쓰는 값 — 줄글을 다시 쪼개지 않고 칸마다 읽는다
    job.person_files, job.common_files = assigned, [str(tmp_path / "공문.pdf")]
    plan = scholarship.plan_person(job, 1)
    assert (plan.student_no, plan.course, plan.common_count) == ("202600010", "운영체제", 1)
    assert plan.files == [job.common_files[0], files[0], files[1]] and plan.summary == job.row_summary(1)


def test_same_name_needs_student_no_in_file_name(tmp_path):
    people = [roster.Person("가나다", "202600001"), roster.Person("가나다", "202600002")]
    files = [str(tmp_path / "가나다.pdf"), str(tmp_path / "가나다_202600002.pdf")]
    assigned, unmatched = scholarship.assign_person_files(files, str(tmp_path), people)
    assert assigned == {1: [files[1]]} and unmatched == [files[0]]


# ── 소득자 팝업 · 첨부 팝업 ─────────────────────────────────────────────
def _grid(*rows):
    return [{"index": i, "text": f"{nm} {no}", "titles": "", "pick": True, "cells": {"k_EMP_NO": no, "k_EMP_NM": nm}}
            for i, (nm, no) in enumerate(rows)]


def test_person_match_uses_student_number_not_name():
    rows = _grid(("가나다", "202600001"), ("가나다", "202500099"))   # 실측: 같은 이름이 두 행
    assert person.match(rows, "가나다", "202600001") == (0, "학번 일치")
    index, why = person.match(rows, "가나다", "")
    assert index is None and "학번이 없어" in why
    index, why = person.match(rows, "가나다", "209999999")
    assert index is None and "없음" in why
    index, why = person.match(_grid(("다른이", "202600001")), "가나다", "202600001")
    assert index is None and "이름이 다름" in why


def test_person_search_tries_student_number_before_name():
    """담당자 요청(2026-10-14): 학번/사번으로 먼저 검색하고, 결과가 없으면 이름으로."""
    assert person.queries_for("가나다", "2026-00001") == [("학번", "202600001"), ("이름", "가나다")]
    assert person.queries_for("가나다", "123456", person.TAB_STAFF) == [("사번", "123456"), ("이름", "가나다")]
    assert person.queries_for("가나다", "") == [("이름", "가나다")]


class FakeGrid:
    """가상 스크롤 결과 목록 — 한 줄 20px, 보이는 곳 근처 21줄만 그린다(실측: 흔한 이름 검색이 늘 '21행')."""
    LINE, RENDER = 20, 21

    def __init__(self, people, visible=10):
        self.people = people
        self.top = 0
        self.height = visible * self.LINE
        self.picked = None

    def _first(self):
        return max(0, self.top // self.LINE - 2)

    def _drawn(self):
        first = self._first()
        return self.people[first:first + self.RENDER]

    def evaluate(self, js, arg=None):
        if js is payee.RESULT_COUNT_JS:
            return {"rows": len(self._drawn()), "kind": "grid"}
        if js is payee.ROWS_JS:
            rows = _grid(*self._drawn())
            for row in rows:
                row["r"] = str(self._first() + row["index"])         # JGM 줄 번호 — 같은 글자의 두 줄도 구분된다
            return {"kind": "grid", "rows": rows}
        if js is person.SCROLL_JS:
            total = len(self.people) * self.LINE
            if arg == "top":
                self.top = 0
            elif isinstance(arg, int):
                self.top = arg
            elif arg == "next":
                self.top += max(40, self.height - 60)
            self.top = max(0, min(self.top, total - self.height))       # 브라우저처럼 끝에서 멈춤
            return {"ok": True, "top": self.top, "h": self.height, "sh": total}
        if js is payee.PICK_ROW_JS:
            self.picked = self._first() + arg
            return {"ok": True}
        raise AssertionError(js[:40])


def _crowd(n, hits):
    return [("가나다", "202600001" if i in hits else str(202500100 + i)) for i in range(n)]


def test_person_scan_scrolls_past_the_drawn_rows():
    """실제 ERP 18:10 — '김형준' 검색이 21행만 읽혀 학번을 못 찾고 멈췄다. 끝까지 내려가 찾는다."""
    grid, logs = FakeGrid(_crowd(30, {25})), []
    index, why, where = person._scan(grid, "가나다", "202600001", logs.append, lambda s: None)
    assert index is not None and why == "학번 일치" and where is grid
    assert grid._drawn()[index] == ("가나다", "202600001")
    assert any("30행 중 학번 202600001 인 행 1개" in line for line in logs)


def test_person_scan_checks_the_whole_list_for_a_second_match():
    """처음 그려진 줄에 하나 보여도, 아래에 같은 학번이 또 있으면 사람에게."""
    index, why, _ = person._scan(FakeGrid(_crowd(30, {3, 27})), "가나다", "202600001", lambda m: None, lambda s: None)
    assert index is None and "2개" in why
    index, why, _ = person._scan(FakeGrid(_crowd(30, set())), "가나다", "202600001", lambda m: None, lambda s: None)
    assert index is None and why == "30행 중 학번 202600001 인 행이 없음"


def test_person_scan_short_list_needs_no_scrolling():
    grid = FakeGrid(_crowd(3, {1}))
    index, why, _ = person._scan(grid, "가나다", "202600001", lambda m: None, lambda s: None)
    assert (index, why, grid.top) == (1, "학번 일치", 0)


def test_attachable_skips_archives_and_temp_files(tmp_path):
    keep = tmp_path / "a.pdf"
    keep.write_bytes(b"%PDF")
    for name in ("b.zip", "~$c.xlsx", "d.lnk"):
        (tmp_path / name).write_bytes(b"x")
    assert attach.attachable([str(tmp_path / n) for n in ("a.pdf", "b.zip", "~$c.xlsx", "d.lnk", "없음.pdf")]) == [str(keep)]


# ── 흐름 (가짜 작성 화면) ──────────────────────────────────────────────
class FakeForm:
    """인건비 탭 작성 화면 — 탭 클릭 · select/text · 저장(청구내역 한 줄 + REQ_CNT)을 흉내."""

    def __init__(self, owner_result="성공:{name}", save_works=True, deduction=0, account=None,
                 reset_income=False, lock_income=False, owner_check="성공:{name}", clear_owner_on_account=False, foreign=()):
        self.v = {"TAB_GB": "10", "APPR_DIV_CD": "", "REQ_CNT": "", "REQ_SEQ_NO": "TMP_REQ_1", "MOD_YN": "N",
                  "OWNR_INQ_RSLT": "", "ACCT_NO": "", "RCV_CD": "", "RCV_NM": "", "fileList": 0}
        self.text = {}
        self.rows = 0
        self.saved = []
        self.clicks = []
        self.owner_result = owner_result
        self.save_works = save_works
        self.deduction = deduction
        self.account = account
        self.reset_income = reset_income      # 소득자를 고르면 ERP 가 소득구분을 기타소득으로 되돌린다(담당자 첫 시험에서 본 것)
        self.lock_income = lock_income        # 소득구분을 바꿔도 그대로 — 사람이 고쳐야 하는 경우
        self.owner_check = owner_check        # [예금주조회] 를 누르면 나오는 결과. None 이면 아무 일도 안 일어난다
        self.clear_owner_on_account = clear_owner_on_account     # 계좌를 고치면 예금주조회 결과가 지워지는 ERP(미실측 — 두 경우 다 본다)
        self.foreign = set(foreign)           # 외국인 학생의 학번 — 외국인등록번호(뒷자리 7) · 영문 예금주
        self.sheet: list[str] = []            # 청구내역 표에 저장된 줄의 글자 (SHEET_ROWS_JS)
        self.total = 0                        # 합계 줄의 청구액

    def pick_person(self, name, no):
        # ERP 가 소득자를 고르면 채우는 계좌 — 기본은 명단(make_ga)과 같은 계좌
        # account="" 는 인적정보에 계좌가 없는 학생 — ERP 가 계좌 · 예금주 칸을 비워 둔다(실제 시험 2026-10-14 19:47)
        acct = self.account if self.account is not None else {"202600001": "12345678901", "202600002": "100000000001"}.get(no, "")
        self.v.update(RCV_CD=no, RCV_NM=name, ACCT_NO=acct, OWNR=name if acct else "", BNK_CD="국민은행" if acct else "",
                      OWNR_INQ_RSLT=self.owner_result.format(name=name),
                      VW_SSN="######-3******", NATION_CD="KR", NATION_NM="대한민국")
        if no in self.foreign:
            self.v.update(VW_SSN="######-7******", OWNR="SUNATULLAEV GO", OWNR_INQ_RSLT="성공:SUNATULLAEV GO")
        if self.reset_income:
            self.v["INCM_DIV_CD"] = self.text["INCM_DIV_CD"] = "기타소득"

    def evaluate(self, js, arg=None):
        if js is form.PROBE_JS:
            return True
        if js is payee.CLICK_ID_JS:
            self.clicks.append(arg)
            if arg == scholarship.TAB_ID:
                self.v["TAB_GB"] = "20"
            # 은행 · 계좌가 비어 있으면 ERP 는 '예금주조회시 은행/계좌번호는 필수입력입니다.' 알림만 띄우고 결과를 안 바꾼다
            if arg == scholarship.OWNER_CHECK_ID and self.owner_check is not None and self.v.get("ACCT_NO"):
                self.v["OWNR_INQ_RSLT"] = self.owner_check.format(name=self.v["RCV_NM"])
            return {"ok": True}
        if js is scholarship.VISIBLE_ID_JS:
            return True
        if js is form.STATE_JS:
            out = {}
            for i in arg:
                info = {"tag": "input", "value": str(self.v.get(i, ""))}
                if i in self.text:
                    info["text"] = self.text[i]
                if i == "fileList":
                    info["options"] = ["f"] * int(self.v["fileList"])
                out[i] = info
            return out
        if js is form.SET_SELECT_JS:
            element_id, wanted = arg
            if element_id == "INCM_DIV_CD" and self.lock_income and self.text.get(element_id):
                return {"ok": True, "now": self.text[element_id]}
            self.v[element_id], self.text[element_id] = wanted, wanted
            return {"ok": True, "now": wanted}
        if js is form.SET_TEXT_JS:
            element_id, value = arg
            if element_id == "ACCT_NO" and self.clear_owner_on_account:
                self.v["OWNR_INQ_RSLT"] = ""
            self.v[element_id] = value
            return {"ok": True, "now": value}
        if js is form.CLICK_SAVE_JS:
            self.clicks.append("저장:" + arg)
            if self.save_works:
                self.saved.append(dict(self.v))
                # 실측 ⑤ 처럼 줄 글자에 적요와 금액(콤마)이 함께 있다
                self.sheet.append("%d %s %s" % (self.rows + 1, self.v.get("REQ_PTCL", ""), f"{int(self.v.get('REQ_AMT') or 0):,}"))
                self.total += int(self.v.get("REQ_AMT") or 0)
                self.rows += 1
                self.v.update(REQ_CNT="322", APPR_DIV_CD="40", REQ_SEQ_NO=f"TMP_REQ_{self.rows + 1}", TAB_GB="10",
                              RCV_CD="", RCV_NM="", OWNR_INQ_RSLT="", ACCT_NO="", REQ_AMT="", REQ_PTCL="")
                self.text.clear()
            return {"ok": True}
        if js is scholarship.SHEET_ROWS_JS:
            return list(self.sheet)
        if js is scholarship.SUMMARY_JS:
            return {"rows": self.rows, "totals": [self.total, self.total - self.deduction, self.deduction] if self.rows else []}
        raise AssertionError("모르는 JS: " + str(js)[:60])


@pytest.fixture
def job(tmp_path):
    folder = tmp_path / "건"
    folder.mkdir()
    make_ga(folder / "명단.xlsx")
    (folder / "공통증빙").mkdir()
    (folder / "공통증빙" / "공문.pdf").write_bytes(b"%PDF")
    return scholarship.load_job(str(folder))


def _patch_popups(monkeypatch, frame, picks=True, attach_ok=True, nation=None):
    """nation: 국적 팝업에서 앱이 고를 (코드, 이름) — None 이면 못 고른다(검색 결과 없음)."""
    calls = {"person": [], "attach": [], "nation": []}

    def fake_nation(page, fr, name, log, wait=None, messages=None):
        calls["nation"].append(name)
        if nation is None:
            return scholarship.nation_mod.NationReport(ok=True, nation=name, note="검색 결과가 없음")
        frame.v.update(NATION_CD=nation[0], NATION_NM=nation[1])
        return scholarship.nation_mod.NationReport(ok=True, nation=name, picked=True, code=nation[0], name=nation[1],
                                                   note="국적 이름 일치")

    monkeypatch.setattr(scholarship.nation_mod, "pick", fake_nation)

    def fake_search(page, fr, name, no, log, wait=None, messages=None, tab=None, auto_pick=True):
        calls["person"].append((name, no, tab))
        if picks:
            frame.pick_person(name, no)
        return person.PersonReport(ok=True, name=name, student_no=no, picked=picks, chosen_no=no if picks else "",
                                   pick_note="학번 일치" if picks else "검색 결과가 없음")

    def fake_upload(page, fr, paths, log, wait=None, messages=None):
        calls["attach"].append(list(paths))
        if attach_ok:
            frame.v["fileList"] = int(frame.v["fileList"]) + len(paths)
        return attach.AttachReport(ok=attach_ok, files=list(paths), uploaded=len(paths) if attach_ok else 0,
                                   note="" if attach_ok else "팝업이 열리지 않았습니다")

    monkeypatch.setattr(scholarship.person, "search", fake_search)
    monkeypatch.setattr(scholarship.attach, "upload", fake_upload)
    return calls


def test_run_puts_each_student_and_saves_without_pressing_apply(monkeypatch, job):
    frame = FakeForm()
    calls = _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and not report.stopped and asked == []
    assert [c[:2] for c in calls["person"]] == [("가나다", "202600001"), ("라마바", "202600002")]
    first, second = frame.saved
    assert first["PAY_EXP_CD"] == "장학금" and first["INCM_DIV_CD"] == "장학금"         # 비목=장학금 → 소득구분=장학금
    assert first["REQ_AMT"] == "300000" and first["REQ_PTCL"] == job.main_summary + "_가나다"
    assert first["APPL_CONT"] == job.main_summary and first["PAY_PLAN_YM"] == job.pay_ym
    assert second["REQ_PTCL"] == job.main_summary + "_라마바"
    assert frame.clicks.count(scholarship.TAB_ID) == 2                            # 저장 뒤 탭이 돌아가므로 사람마다 다시
    assert all(c in (scholarship.TAB_ID, "저장:저장") for c in frame.clicks)     # 누른 것은 탭과 [저장] 뿐
    assert calls["attach"][0] == job.files_for(0)
    assert "[신청]" in report.lines()[-1]
    # 넣기를 마치면 청구내역 표를 다시 읽어 명단과 맞춘다(신청 전 대조) — 읽기만 하므로 누른 것은 그대로
    assert report.check is not None and report.check.ok, report.check.problems
    assert "신청 전 대조: 청구내역 2줄 — 명단과 맞습니다" in report.lines()


def test_income_kind_is_picked_again_when_erp_resets_it_after_person(monkeypatch, job):
    """담당자 첫 시험(2026-10-14): 소득구분을 '장학금' 으로 골랐는데 저장 전에는 다른 값 — 소득자를 고른 뒤 다시 고른다."""
    frame = FakeForm(reset_income=True)
    _patch_popups(monkeypatch, frame)
    logs, asked = [], []
    report = scholarship.run(object(), frame, job, logs.append, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and asked == []
    assert all(row["INCM_DIV_CD"] == "장학금" for row in frame.saved)
    assert any("다시 고름" in line for line in report.results[0].lines)
    assert any("기타소득" in line and "다시" in line for line in logs)            # 무엇으로 바뀌어 있었는지 로그에 남긴다


def test_income_kind_that_will_not_change_is_saved_and_listed_for_review(monkeypatch, job):
    """담당자 결정(2026-10-14): 앱이 다시 골라도 소득구분이 안 바뀌면 멈추지 말고 그대로 임시저장 + 신청 전 확인 목록."""
    frame = FakeForm(reset_income=True, lock_income=True, deduction=26400)
    _patch_popups(monkeypatch, frame)
    asked, logs, notified = [], [], []
    report = scholarship.run(object(), frame, job, logs.append, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False, on_person=lambda i, st, n: notified.append((i, st)))
    assert report.saved == 2 and not report.stopped and asked == []             # 묻지도 멈추지도 않는다
    assert [row["INCM_DIV_CD"] for row in frame.saved] == ["기타소득", "기타소득"]   # 그대로 저장
    review = report.review_items()
    assert len(review) == 3 and all(item.startswith(("1. 가나다", "2. 라마바")) for item in review)
    assert "기타소득" in review[0] and "[신청]" in review[0]
    assert "공제 26,400원" in review[1]                                         # 첫 줄을 넣고 공제가 생겼다
    assert not any("공제" in item for item in report.results[1].review)         # 둘째 줄은 공제가 늘지 않았다(합계 그대로)
    assert ("review" in [st for _i, st in notified]) and "신청 전 확인" in "\n".join(report.lines())
    assert any("신청 전 확인" in line for line in logs)


def test_owner_check_is_pressed_by_the_app_when_result_is_empty(monkeypatch, job):
    """담당자 결정(2026-10-14): 예금주조회는 앱이 누른다 — '성공:학생이름' 이 나오면 묻지 않고 이어 간다."""
    frame = FakeForm(owner_result="")
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and asked == [] and frame.clicks.count(scholarship.OWNER_CHECK_ID) == 2
    assert any("앱이 누름" in line for line in report.results[0].lines)


def test_owner_check_for_someone_else_goes_to_a_person(monkeypatch, job):
    frame = FakeForm(owner_result="", owner_check="성공:다른사람")   # 계좌 주인이 학생이 아니다
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 0 and report.stopped and "다른사람" in asked[0]


def test_foreign_student_is_told_by_registration_number_then_english_owner():
    assert scholarship.foreign_reason("######-7******", ["성공:SUNATULLAEV GO"]) == "외국인등록번호"
    assert scholarship.foreign_reason("######-3******", ["성공:SUNATULLAEV GO"]) == ""      # 내국인 — 영문 예금주만으로는 아니다
    assert scholarship.foreign_reason("", ["", "성공:SUNATULLAEV GO"]) == "예금주가 영문 이름(SUNATULLAEV GO)"
    assert scholarship.foreign_reason("", ["가나다", "성공:가나다"]) == ""
    assert scholarship.latin_name("SUNATULLAEV GO") == "SUNATULLAEV GO" and scholarship.latin_name("가나다 GO") == ""


def test_english_owner_passes_when_success():
    """실제 시험 18:25 — 외국인 학생의 예금주조회 결과 '성공:SUNATULLAEV GO' 를 한글 이름과 달라 사람에게 넘겼다."""
    assert scholarship._owner_result_ok("성공:SUNATULLAEV GO", "수나툴라예브 골리브존") == (True, "성공:SUNATULLAEV GO")
    assert not scholarship._owner_result_ok("성공:다른사람", "가나다")[0]
    assert not scholarship._owner_result_ok("실패:SUNATULLAEV GO", "가나다")[0]


def test_foreign_student_asks_for_nationality_then_saves_with_review(monkeypatch, job):
    """담당자 2026-10-14: 외국인 학생은 개별증빙의 신분증을 보고 국적을 고른다 — 국적 팝업은 미실측이라 사람이, 앱은 KR 이 아닌지만 본다."""
    frame = FakeForm(foreign={"202600002"})
    _patch_popups(monkeypatch, frame)
    asked = []

    def ask(message):
        asked.append(message)
        frame.v.update(NATION_CD="UZ", NATION_NM="우즈베키스탄")     # 사람이 국적 돋보기로 고름
        return True

    logs = []
    report = scholarship.run(object(), frame, job, logs.append, ask=ask, wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and not report.stopped
    assert len(asked) == 1 and "라마바" in asked[0] and "외국인" in asked[0] and "신분증" in asked[0]
    assert frame.saved[1]["NATION_CD"] == "UZ" and frame.saved[0]["NATION_CD"] == "KR"
    assert scholarship.OWNER_CHECK_ID not in frame.clicks             # 영문 예금주 '성공' 은 다시 조회하지 않는다
    review = report.review_items()
    assert len(review) == 1 and "라마바" in review[0] and "SUNATULLAEV GO" in review[0]
    assert any("외국인 학생으로 봄: 2. 라마바 (외국인등록번호)" in line for line in logs)
    assert not any("######" in line for line in logs)                  # 주민번호 칸 값은 로그에 남기지 않는다


def _nation_job(tmp_path, nation):
    folder = tmp_path / "국적"
    folder.mkdir()
    add_nation(make_ga(folder / "명단.xlsx"), {6: nation})
    return scholarship.load_job(str(folder))


def test_roster_nation_marks_foreign_student_even_with_korean_looking_number(monkeypatch, tmp_path):
    """명단에 국적을 적으면 주민번호 칸이 내국인처럼 보여도 외국인으로 보고, 띠에 그 나라를 보여 준다."""
    job = _nation_job(tmp_path, "우즈베키스탄")
    plan = scholarship.plan_person(job, 1)
    assert "외국인(우즈베키스탄)" in plan.note and any("국적 우즈베키스탄" in line for line in plan.lines)
    assert not scholarship.plan_person(job, 0).note
    frame = FakeForm()                                           # 주민번호 칸은 '-3' (내국인처럼)
    calls = _patch_popups(monkeypatch, frame, nation=("UZ", "우즈베키스탄"))
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or True,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and asked == []                     # 앱이 국적조회에서 골라 사람에게 묻지 않는다
    assert calls["nation"] == ["우즈베키스탄"]                     # 외국인 학생 한 명만
    assert frame.saved[1]["NATION_CD"] == "UZ" and report.review_items() == []
    assert any("국적: 우즈베키스탄(UZ)" in line for line in report.results[1].lines)


def test_nation_falls_back_to_person_when_app_cannot_pick(monkeypatch, tmp_path):
    job = _nation_job(tmp_path, "우즈베키스탄")
    frame = FakeForm()
    _patch_popups(monkeypatch, frame)                              # 국적 팝업에서 못 고름
    asked = []

    def ask(message):
        asked.append(message)
        frame.v.update(NATION_CD="UZ", NATION_NM="우즈베키스탄 공화국")
        return True

    report = scholarship.run(object(), frame, job, lambda m: None, ask=ask, wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and len(asked) == 1
    assert "명단 국적: 우즈베키스탄" in asked[0] and "앱이 고르지 못함: 검색 결과가 없음" in asked[0]
    assert report.review_items() == []                            # 사람이 고른 국적이 명단과 같다


def test_nation_is_not_picked_when_auto_pick_is_off(monkeypatch, tmp_path):
    job = _nation_job(tmp_path, "우즈베키스탄")
    frame = FakeForm()
    calls = _patch_popups(monkeypatch, frame, nation=("UZ", "우즈베키스탄"))
    asked = []

    def ask(message):
        asked.append(message)
        frame.v.update(NATION_CD="UZ", NATION_NM="우즈베키스탄")
        return True

    report = scholarship.run(object(), frame, job, lambda m: None, ask=ask, wait=lambda s: None, dry_run=False, auto_pick=False)
    assert report.saved == 2 and calls["nation"] == [] and len(asked) == 1


def test_nation_different_from_roster_is_saved_with_review(monkeypatch, tmp_path):
    job = _nation_job(tmp_path, "베트남")
    frame = FakeForm()
    _patch_popups(monkeypatch, frame)

    def ask(message):
        frame.v.update(NATION_CD="CN", NATION_NM="중국")
        return True

    report = scholarship.run(object(), frame, job, lambda m: None, ask=ask, wait=lambda s: None, dry_run=False)
    review = report.review_items()
    assert report.saved == 2 and len(review) == 1 and "중국" in review[0] and "베트남" in review[0]


def test_foreign_student_stops_when_nationality_is_left_korean(monkeypatch, job):
    frame = FakeForm(foreign={"202600001"})
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 0 and report.stopped and frame.saved == []
    assert "국적" in asked[0] and "대한민국" in asked[0]


def test_run_refuses_to_mix_folders_into_another_sheet(monkeypatch, job):
    """담당자 2026-10-14: 폴더가 여럿이면 하나를 끝내고 다음으로 — 앞 폴더의 청구서에 다음 폴더 사람이 섞이면 안 된다."""
    frame = FakeForm()
    frame.v["REQ_CNT"], frame.sheet = "500", ["02_2026년도 8월 (다)형 근로장학금_홍길동"]     # 앞 폴더를 넣은 청구서가 떠 있다
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 0 and report.stopped and frame.saved == [] and "다른 청구서" in report.blocked
    assert "새 청구서" in asked[0] and "결의번호 500" in asked[0]


def test_run_continues_after_person_opens_a_new_sheet(monkeypatch, job):
    frame = FakeForm()
    frame.v["REQ_CNT"], frame.sheet = "500", ["02_다른 청구서_홍길동"]
    _patch_popups(monkeypatch, frame)

    def ask(message):
        frame.v["REQ_CNT"], frame.sheet = "", []                  # 사람이 [새 청구서] 를 띄움
        return True

    report = scholarship.run(object(), frame, job, lambda m: None, ask=ask, wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and job.req_cnt == "322"              # 첫 [저장] 뒤 결의번호를 기억한다


def test_run_continues_in_this_folders_own_sheet(monkeypatch, job):
    """멈춘 뒤 남은 사람만 다시 넣을 때 — 화면의 청구서가 이 폴더 것이면(적요가 대표적요로 시작) 묻지 않고 이어 간다."""
    frame = FakeForm()
    frame.v["REQ_CNT"], frame.sheet = "777", [job.row_summary(0)]
    job.selected = [False, True]
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and asked == []
    other = scholarship.load_job(job.folder)
    assert scholarship.sheet_owner(other, "", []) == (True, "새 청구서")
    assert scholarship.sheet_owner(other, "900", ["03_다른 대표적요_가나다"])[0] is False
    # 결의번호가 없으면 아직 저장 전인 새 청구서다 — 표에 줄이 보여도(새 청구서의 빈 줄, 담당자 실제 시험 2026-10-19)
    assert scholarship.sheet_owner(other, "", ["1"]) == (True, "새 청구서")
    # 결의번호를 못 읽었어도 진행구분이 임시저장(40)이면 저장된 청구서로 본다
    assert scholarship.sheet_owner(other, "", ["03_다른 대표적요_가나다"], "40")[0] is False


def test_other_saved_sheet_is_replaced_by_a_new_one_the_app_opens(job):
    """다른 폴더의 저장된 청구서가 떠 있으면 open_sheet(결의서 목록의 [일반청구])로 새 청구서를 띄우고 묻지 않는다(담당자 허락 2026-10-19).
    띄우지 못하면 전처럼 사람에게 묻는다."""
    class Sheet:
        def __init__(self, req_cnt, rows):
            self.req_cnt, self.rows = req_cnt, rows

        def evaluate(self, js, arg=None):
            if js is form.PROBE_JS:
                return True
            if js is scholarship.SHEET_ROWS_JS:
                return self.rows
            if js is form.STATE_JS:
                return {k: {"value": {"REQ_CNT": self.req_cnt, "APPR_DIV_CD": "40" if self.req_cnt else ""}.get(k, "")} for k in arg}
            raise AssertionError(str(js)[:40])

    other, fresh = Sheet("900", ["03_다른 대표적요_가나다"]), Sheet("", [])
    asked, logs = [], []
    ok, frame, why = scholarship._ensure_own_sheet(object(), other, job, lambda m: asked.append(m) or False, lambda s: None,
                                                   logs.append, open_sheet=lambda: (True, "새 청구서", fresh))
    assert ok and frame is fresh and why == "새 청구서" and asked == []
    ok, frame, why = scholarship._ensure_own_sheet(object(), other, job, lambda m: asked.append(m) or False, lambda s: None,
                                                   logs.append, open_sheet=lambda: (False, "결의서 목록이 열려 있지 않음", None))
    assert not ok and frame is other and len(asked) == 1 and any("결의서 목록이 열려 있지 않음" in line for line in logs)


def test_sheet_rows_skip_the_empty_row_of_a_new_sheet():
    """새 청구서의 청구내역 표에는 선택 칸만 있는 빈 줄이 있다 — 저장된 줄로 세지 않는다(실제 Chromium 이 있을 때만)."""
    sync_api = pytest.importorskip("playwright.sync_api")
    html = ("<table><tr><th>○</th><th>NO</th><th>적요</th></tr>"
            "<tr><td><input type=radio name=CUR_REQ></td><td></td><td></td></tr>"
            "<tr><td><input type=radio name=CUR_REQ></td><td>1</td><td></td></tr>"
            "<tr><td><input type=radio name=CUR_REQ></td><td colspan=2>조회된 데이터가 없습니다.</td></tr>"
            "<tr><td><input type=radio name=CUR_REQ></td><td>2</td><td>03_진행요원비_가나다 240,000</td></tr>"
            "<tr><td><input type=radio name=CUR_REQ></td><td>3</td><td title='03_진행요원비_라마바'>03_진행요원비_ ...</td></tr></table>")
    try:
        with sync_api.sync_playwright() as p:
            path = "/opt/pw-browsers/chromium"
            browser = p.chromium.launch(executable_path=path if os.path.exists(path) else None)
            page = browser.new_page()
            page.set_content(html)
            rows = page.evaluate(scholarship.SHEET_ROWS_JS)
            browser.close()
    except Exception as error:  # noqa: BLE001 — 브라우저가 없는 곳(CI)
        pytest.skip(f"브라우저 없음: {str(error).splitlines()[0][:80]}")
    # 잘려 보이는 칸의 title(전체 글)은 줄 글자 뒤에 붙는다
    assert rows == ["2 03_진행요원비_가나다 240,000", "3 03_진행요원비_ ... 03_진행요원비_라마바"]


def test_find_form_prefers_the_visible_sheet():
    """앞 폴더의 청구서 탭이 남은 채 새 청구서를 띄우면 작성 화면이 둘 — 보이는 쪽을 고른다."""
    class Frame:
        def __init__(self, visible):
            self.visible = visible

        def evaluate(self, js, arg=None):
            if js is form.PROBE_JS:
                return True
            if js is form.FRAME_VISIBLE_JS:
                return self.visible
            raise AssertionError(js[:30])

    hidden, shown = Frame(False), Frame(True)

    class Page:
        frames = [hidden, shown]

    class Context:
        pages = [Page()]

    assert form.find_form(Context())[1] is shown


def test_wait_for_form_retries_when_the_first_look_fails():
    """첫 evaluate 가 False 여도(작성 화면을 띄워 뒀는데 '찾지 못했습니다', 담당자 실제 시험 2026-10-19) 다시 보고, 창을 앞으로 가져온다."""
    class Frame:
        url = "https://erp/rexpe_0084_01.act?PRJ_NO=1"

        def __init__(self):
            self.looks = 0

        def evaluate(self, js, arg=None):
            if js is form.PROBE_JS:
                self.looks += 1
                return self.looks >= 3                       # 두 번은 못 읽다가 세 번째에
            return True

    frame = Frame()

    class Page:
        frames = [frame]
        fronted = 0
        waited = 0

        def bring_to_front(self):
            Page.fronted += 1

        def wait_for_timeout(self, ms):
            Page.waited += 1

    class Context:
        pages = [Page()]

    logs = []
    page, found = form.wait_for_form(Context(), log=logs.append)
    assert found is frame and Page.fronted == 1 and Page.waited == 2
    assert logs == ["작성 화면을 3번째에 찾았습니다"]
    # 끝내 없으면 None 과 진단 한 줄(주소의 물음표 뒤는 적지 않는다)
    frame.looks, logs[:] = -100, []
    assert form.wait_for_form(Context(), tries=3, log=logs.append) == (None, None)
    assert logs and "작성 화면 주소 프레임 1개" in logs[0] and "PRJ_NO" not in logs[0]


def test_run_stops_and_asks_when_unsure(monkeypatch, job):
    frame = FakeForm(owner_result="", owner_check=None)          # 예금주조회 결과가 비어 있고 눌러도 안 나온다
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 0 and report.stopped and frame.saved == []
    assert asked and "예금주조회" in asked[0]


def test_run_asks_for_person_when_popup_is_ambiguous_then_continues(monkeypatch, job):
    frame = FakeForm()
    _patch_popups(monkeypatch, frame, picks=False)
    answers = []

    def ask(message):
        answers.append(message)
        # 사람이 팝업에서 골랐다 — 작성 화면에 그 학생이 들어온다
        name = "가나다" if "가나다" in message else "라마바"
        frame.pick_person(name, "202600001" if name == "가나다" else "202600002")
        return True

    report = scholarship.run(object(), frame, job, lambda m: None, ask=ask, wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and len(answers) == 2 and "소득자" in answers[0]


def test_run_refuses_to_save_on_account_mismatch_if_person_says_stop(monkeypatch, job):
    frame = FakeForm(account="99999999999")
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 0 and "계좌" in asked[0] and frame.saved == []


def _text_pdf(path, *lines):
    """글자가 든 아주 작은 PDF — 학생이 낸 '지출자료' 처럼 계좌가 적힌 개별증빙을 흉내 낸다(숫자·영문만)."""
    body = "BT /F1 12 Tf 50 750 Td " + " ".join(f"({line}) Tj 0 -16 Td" for line in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(body)} >>\nstream\n{body}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = "%PDF-1.4\n", []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    pathlib.Path(path).write_bytes(out.encode("latin-1"))
    return str(path)


def test_accounts_are_read_from_person_pdfs(tmp_path):
    pdf = _text_pdf(tmp_path / "a.pdf", "Bank 110-594-395796", "Phone 010-1234-5678", "ID 202416354")
    found, read = scholarship.accounts_in_files([pdf, str(tmp_path / "없음.pdf"), str(tmp_path / "b.jpg")])
    assert read == 1 and "110594395796" in found and "202416354" not in found      # 학번(9자리)은 계좌로 보지 않는다


def _mismatch_job(tmp_path, pdf_lines=None):
    folder = tmp_path / "건"
    folder.mkdir()
    make_ga(folder / "명단.xlsx", rows=[("가나다", 202600001, "신한", "110-594-395796", 12000, "25", 300000)])
    if pdf_lines is not None:
        (folder / "개별증빙").mkdir()
        _text_pdf(folder / "개별증빙" / "가나다 지출자료.pdf", *pdf_lines)
    return scholarship.load_job(str(folder))


def test_student_numbers_are_read_from_person_pdfs(tmp_path):
    pdf = _text_pdf(tmp_path / "a.pdf", "ID 202600009", "Again 202600009", "Bank 110-594-395796", "Date 20261019", "Other 202600001")
    numbers, read = scholarship.student_numbers_in_files([pdf])
    assert read == 1 and numbers == ["202600009", "202600001"]           # 많이 나온 것부터 · 날짜(8자리) · 계좌는 아님


def test_person_not_found_by_roster_number_is_rechecked_with_document_number(monkeypatch, tmp_path):
    """명단 학번으로 소득자를 못 찾으면 개별증빙의 학번으로 다시 찾는다(담당자 요청 2026-10-19). 고르면 신청 전 확인에 남긴다."""
    job = _mismatch_job(tmp_path, ["Student 202600009"])
    frame = FakeForm(account="110594395796")
    _patch_popups(monkeypatch, frame)
    searched = []

    def fake_search(page, fr, name, no, log, wait=None, messages=None, tab=None, auto_pick=True):
        searched.append(no)
        if no == "202600009":                                          # ERP 에 있는 진짜 학번 — 이름도 같다
            frame.pick_person(name, no)
            return person.PersonReport(ok=True, name=name, student_no=no, picked=True, chosen_no=no, pick_note="학번 일치")
        return person.PersonReport(ok=True, name=name, student_no=no, rows=1, pick_note=f"1행 중 학번 {no} 인 행이 없음")

    monkeypatch.setattr(person, "search", fake_search)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert searched == ["202600001", "202600009"] and asked == []
    assert report.saved == 1 and frame.saved[0]["RCV_CD"] == "202600009"
    assert any("개별증빙의 학번 202600009 로 골랐습니다" in item for item in report.review_items())


def test_person_recheck_skips_numbers_of_other_students_and_asks_when_nothing_fits(monkeypatch, tmp_path):
    """개별증빙에 명단의 다른 학생 학번만 있으면 다시 찾지 않고 사람에게 — 여럿 · 이름 다름이면 처음부터 재확인하지 않는다."""
    folder = tmp_path / "건"
    folder.mkdir()
    make_ga(folder / "명단.xlsx")
    (folder / "개별증빙").mkdir()
    _text_pdf(folder / "개별증빙" / "가나다 신청서.pdf", "Member 202600002")
    job = scholarship.load_job(str(folder))
    job.selected = [True, False]
    frame = FakeForm()
    _patch_popups(monkeypatch, frame)
    searched = []

    def fake_search(page, fr, name, no, log, wait=None, messages=None, tab=None, auto_pick=True):
        searched.append(no)
        return person.PersonReport(ok=True, name=name, student_no=no, rows=1, pick_note=f"1행 중 학번 {no} 인 행이 없음")

    monkeypatch.setattr(person, "search", fake_search)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert searched == ["202600001"] and report.saved == 0 and asked and "소득자" in asked[0]


def test_account_mismatch_puts_the_document_account_then_presses_owner_check(monkeypatch, tmp_path):
    """담당자 결정(2026-10-14): ERP 계좌가 다르면 첨부서류의 계좌를 넣고, 앱이 [예금주조회] 를 눌러 '성공' 이면 이어 간다."""
    job = _mismatch_job(tmp_path, ["Account 110-594-395796"])
    frame = FakeForm(account="30120104410673", clear_owner_on_account=True)
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or True,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and frame.saved[0]["ACCT_NO"] == "110594395796" and frame.saved[0]["BNK_CD"] == "신한"
    assert asked == [] and scholarship.OWNER_CHECK_ID in frame.clicks
    assert report.review_items() == []                                              # 첨부서류로 확인된 계좌
    assert any("첨부서류의" in line and "넣음" in line and "ERP 에 있던 계좌" in line for line in report.results[0].lines)


def test_owner_check_that_cannot_be_told_fresh_goes_to_a_person(monkeypatch, tmp_path):
    """계좌를 바꿨는데 [예금주조회] 뒤 결과 글이 그대로이고 알림도 없으면 — 남은 '성공' 이 옛 계좌 것일 수 있어 사람에게."""
    job = _mismatch_job(tmp_path, ["Account 110-594-395796"])
    frame = FakeForm(account="30120104410673")                     # 결과 칸이 '성공:가나다' 그대로 남는 ERP
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or True,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and len(asked) == 1 and "바뀌지 않음" in asked[0]


def test_account_mismatch_without_readable_document_uses_roster_and_lists_it(monkeypatch, tmp_path):
    job = _mismatch_job(tmp_path, pdf_lines=None)
    frame = FakeForm(account="30120104410673")
    _patch_popups(monkeypatch, frame)
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: True, wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and frame.saved[0]["ACCT_NO"] == "110594395796"
    assert any("첨부서류에서 계좌를 읽지 못해" in item for item in report.review_items())


def test_account_mismatch_keeps_erp_account_when_document_agrees_with_erp(monkeypatch, tmp_path):
    job = _mismatch_job(tmp_path, ["Account 301-2010-4410-673"])
    frame = FakeForm(account="30120104410673")
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or True,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and frame.saved[0]["ACCT_NO"] == "30120104410673" and asked == []
    assert any("명단 엑셀을 확인" in line for line in report.results[0].lines)


def test_missing_erp_account_is_filled_from_roster_then_owner_check(monkeypatch, tmp_path):
    """실제 시험(2026-10-14 19:47): 인적정보에 계좌가 없는 학생 — 앱이 빈 계좌로 [예금주조회] 를 눌러 멈췄다.
    담당자: 계좌가 없거나 다르면 지급내역(명단) · 개별증빙의 계좌를 넣는다. 넣은 뒤 [예금주조회] 를 눌러 '성공' 이면 이어 간다."""
    job = _mismatch_job(tmp_path, pdf_lines=None)                  # 개별증빙 없음(스캔 PDF 와 같음)
    frame = FakeForm(account="", owner_result="")
    _patch_popups(monkeypatch, frame)
    asked, logs = [], []
    report = scholarship.run(object(), frame, job, logs.append, ask=lambda m: asked.append(m) or True,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and asked == []
    saved = frame.saved[0]
    assert saved["ACCT_NO"] == "110594395796" and saved["BNK_CD"] == "신한" and saved["OWNR_INQ_RSLT"] == "성공:가나다"
    assert saved["OWNR"] == "가나다"                                 # 비어 있던 예금주 칸은 예금주조회가 돌려준 이름으로
    assert frame.clicks.index(scholarship.OWNER_CHECK_ID) >= 0
    assert report.review_items() == []                              # 읽을 첨부서류가 없으면 매번 묻지 않는다
    assert any("ERP 계좌 없음" in line for line in logs)
    assert any("ERP 계좌 칸이 비어 있었음" in line for line in report.results[0].lines)


def test_missing_erp_account_with_document_disagreeing_is_listed(monkeypatch, tmp_path):
    job = _mismatch_job(tmp_path, ["Account 333-3333-3333"])        # 첨부서류에 다른 계좌
    frame = FakeForm(account="", owner_result="")
    _patch_popups(monkeypatch, frame)
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: True, wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and frame.saved[0]["ACCT_NO"] == "110594395796"
    assert any("첨부서류에서 계좌를 읽지 못해" in item for item in report.review_items())


def test_missing_erp_account_without_roster_account_goes_to_a_person(monkeypatch, tmp_path):
    folder = tmp_path / "건"
    folder.mkdir()
    make_ga(folder / "명단.xlsx", rows=[("가나다", 202600001, "신한", "", 12000, "25", 300000)])
    (folder / "개별증빙").mkdir()
    _text_pdf(folder / "개별증빙" / "가나다 통장.pdf", "Account 110-594-395796")
    job = scholarship.load_job(str(folder))
    frame = FakeForm(account="", owner_result="")
    _patch_popups(monkeypatch, frame)
    asked = []

    def ask(message):
        asked.append(message)
        frame.v.update(BNK_CD="신한", ACCT_NO="110594395796", OWNR_INQ_RSLT="성공:가나다")    # 사람이 넣고 조회함
        return True

    report = scholarship.run(object(), frame, job, lambda m: None, ask=ask, wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and len(asked) == 1
    assert "계좌가 없는데" in asked[0] and "명단에 계좌가" in asked[0] and "110594395796" in asked[0]


def test_missing_erp_bank_is_filled_from_roster(monkeypatch, tmp_path):
    """계좌는 있는데 은행이 '해당없음' 이어도 예금주조회는 '은행/계좌번호는 필수' 를 낸다 — 명단의 은행을 넣고 조회한다."""
    job = _mismatch_job(tmp_path, pdf_lines=None)

    class NoBank(FakeForm):
        def pick_person(self, name, no):
            super().pick_person(name, no)
            self.v["BNK_CD"] = ""

    frame = NoBank(account="110594395796", owner_result="")
    _patch_popups(monkeypatch, frame)
    logs = []
    report = scholarship.run(object(), frame, job, logs.append, ask=lambda m: True, wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and frame.saved[0]["BNK_CD"] == "신한" and frame.saved[0]["OWNR_INQ_RSLT"] == "성공:가나다"
    assert report.review_items() == []
    assert any("ERP 은행 없음" in line for line in logs)
    assert any(line.startswith("소득자를 고른 뒤 ERP 계좌: 은행 (비어 있음) · 계좌 110594395796") for line in logs)


def test_erp_account_filled_late_is_waited_for(monkeypatch, tmp_path):
    """소득자를 고른 뒤 ERP 가 계좌를 조금 늦게 채우면 — 비었다고 보고 덮어쓰지 않고 기다린다."""
    job = _mismatch_job(tmp_path, pdf_lines=None)

    class Late(FakeForm):
        reads = 0

        def pick_person(self, name, no):
            super().pick_person(name, no)
            self.v["ACCT_NO"] = ""
            self.reads = 0

        def evaluate(self, js, arg=None):
            if js is form.STATE_JS and "ACCT_NO" in arg and self.v.get("RCV_CD") and not self.v["ACCT_NO"]:
                self.reads += 1
                if self.reads == 3:
                    self.v["ACCT_NO"] = "110-594-395796"
            return super().evaluate(js, arg)

    frame = Late()
    _patch_popups(monkeypatch, frame)
    logs = []
    report = scholarship.run(object(), frame, job, logs.append, ask=lambda m: True, wait=lambda s: None, dry_run=False)
    assert report.saved == 1 and not any("ERP 계좌 없음" in line for line in logs)
    assert scholarship.OWNER_CHECK_ID not in frame.clicks


def test_nation_written_as_code_is_compared_by_code(monkeypatch, tmp_path):
    """실제 시험(2026-10-14): 명단 국적 'UZ' 와 ERP '우즈베키스탄(UZ)' 를 이름으로만 맞춰 '다름' 이라고 잘못 남겼다."""
    assert scholarship.nation_mod.same_nation("UZ", "uz", "우즈베키스탄")
    assert not scholarship.nation_mod.same_nation("VN", "UZ", "우즈베키스탄")
    assert scholarship.nation_mod.same_nation("우즈베키스탄", "UZ", "우즈베키스탄 공화국")
    job = _nation_job(tmp_path, "UZ")
    frame = FakeForm()
    _patch_popups(monkeypatch, frame, nation=("UZ", "우즈베키스탄"))
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: True, wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and report.review_items() == []


def test_run_flags_nonzero_deduction_after_save(monkeypatch, job):
    """공제가 0 이 아니면 멈추지 않고 신청 전 확인 목록으로(담당자 결정 2026-10-14) — 늘어난 줄에만 붙는다."""
    frame = FakeForm(deduction=26400)
    _patch_popups(monkeypatch, frame)
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and not report.stopped and asked == [] and len(frame.saved) == 2
    assert report.review_items() == [report.review_items()[0]] and "공제 26,400원" in report.review_items()[0]


def test_dry_run_touches_nothing(job):
    report = scholarship.run(None, None, job, lambda m: None, ask=lambda m: False, dry_run=True)
    assert report.dry_run and [r.status for r in report.results] == ["planned", "planned"]
    assert "가나다" in "\n".join(report.lines()) and report.lines()[0].startswith("미리 보기")


def test_save_detection_follows_measured_reload():
    """실측 ⑤: 저장 뒤 REQ_CNT 가 생기고 APPR_DIV_CD=40, MOD_YN 은 N 그대로."""
    before = {"REQ_SEQ_NO": "TMP_REQ_1", "MOD_YN": "N", "APPR_DIV_CD": "", "REQ_CNT": ""}
    assert form._changed(before, dict(before, REQ_CNT="322", APPR_DIV_CD="40", REQ_SEQ_NO="TMP_REQ_2"))
    assert not form._changed(before, dict(before))


def test_previous_students_save_message_does_not_count_for_the_next(monkeypatch, job):
    """둘째 학생의 [저장] 이 안 먹었는데 첫 학생의 '저장되었습니다' 가 남아 있어 끝난 것으로 보면 안 된다."""
    frame = FakeForm()
    _patch_popups(monkeypatch, frame)
    live = []
    original = frame.evaluate

    def evaluate(js, arg=None):
        if js is form.CLICK_SAVE_JS:
            if frame.rows == 0:
                live.append("저장되었습니다.")
                return original(js, arg)
            frame.clicks.append("저장:" + arg)          # 둘째 학생: 눌렀지만 아무 일도 없음
            return {"ok": True}
        return original(js, arg)

    frame.evaluate = evaluate
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, messages=live, dry_run=False)
    assert report.saved == 1 and report.stopped and report.results[1].status == "stopped"
    assert report.messages == ["저장되었습니다."]


# ── 청구 종류 (일반청구 화면에서 폴더마다 고른다) ─────────────────────────
def test_kind_guess_recognizes_scholarship_folder_shape(tmp_path):
    from claim_downloader.claims import kinds

    general = tmp_path / "정수기 9월"
    general.mkdir()
    (general / "세금계산서.pdf").write_bytes(b"%PDF")
    assert kinds.guess(str(general)) == kinds.GENERAL

    sch_dir = tmp_path / "7월 (가)형"
    sch_dir.mkdir()
    make_ga(sch_dir / "명단.xlsx")
    assert kinds.guess(str(sch_dir)) == kinds.GENERAL           # 엑셀만 있고 증빙 폴더가 없으면 아직 일반
    (sch_dir / "개별증빙").mkdir()
    assert kinds.guess(str(sch_dir)) == kinds.SCHOLARSHIP
    (sch_dir / "~$명단.xlsx").write_bytes(b"x")                  # 엑셀이 열려 있을 때 생기는 임시 파일은 엑셀로 치지 않는다
    assert kinds.looks_like_scholarship(str(sch_dir))


def test_kind_guess_reads_payment_excel_without_evidence_folders(tmp_path):
    """공통/개별 폴더가 없어도 '지급내역' 엑셀 + 장학금 · 혁신인재지원금 이름이면 명단 청구 — AI 가 PDF 를 읽지 않게(2026-10-19)."""
    from claim_downloader.claims import kinds

    def folder(name, excel):
        d = tmp_path / name
        d.mkdir(parents=True)
        (d / excel).write_bytes(b"x")
        (d / (excel.rsplit(".", 1)[0] + ".pdf")).write_bytes(b"%PDF")
        return str(d)

    assert kinds.guess(folder("9월", "2026학년도 (나)형 근로장학금 지급내역.xlsx")) == kinds.SCHOLARSHIP
    assert kinds.guess(folder("소학회/9월", "진행요원비(소학회 회장) 혁신인재지원금 지급내역.xlsx")) == kinds.INNOVATION
    # 이름에 그 말이 없는 지급내역(자문료) · '지원내역'(프로그램참여지원비)은 일반청구라 그대로
    assert kinds.guess(folder("전문가/자문", "자문료 지급내역.xlsx")) == kinds.GENERAL
    assert kinds.guess(folder("학회", "(양식) 프로그램참여지원비 지원내역.xlsx")) == kinds.GENERAL
    assert kinds.roster_excels(str(tmp_path / "9월")) == ["2026학년도 (나)형 근로장학금 지급내역.xlsx"]


def test_kind_resolve_honors_choice_and_refuses_unready(tmp_path):
    from claim_downloader.claims import kinds

    folder = tmp_path / "건"
    folder.mkdir()
    assert kinds.resolve(kinds.AUTO, str(folder)) == kinds.GENERAL
    assert kinds.resolve(kinds.SCHOLARSHIP, str(folder)) == kinds.SCHOLARSHIP     # 담당자가 고르면 모양과 상관없이 그 종류
    for key in (kinds.TRAVEL, kinds.CARD):
        assert not kinds.find(key).ready
        with pytest.raises(ValueError):
            kinds.resolve(key, str(folder))
    assert [k.name for k in kinds.KINDS] == ["일반", "장학금", "혁신인재지원금", "출장비", "카드"]
    assert kinds.resolve(kinds.INNOVATION, str(folder)) == kinds.INNOVATION and kinds.uses_roster(kinds.INNOVATION)


def _roster_folder(base, *parts, excel="명단.xlsx"):
    folder = base.joinpath(*parts)
    (folder / "공통증빙").mkdir(parents=True)
    make_ga(folder / excel)
    return folder


def test_auto_kind_tells_innovation_from_scholarship_by_names(tmp_path):
    """담당자 2026-10-14: 혁신인재지원금은 장학금과 같은 폴더 모양 · 같은 흐름이고 비목만 다르다. '자동' 은 이름으로 가른다 —
    둘을 함께 담은 묶음 폴더('4. 장학금 및 혁신인재지원금')는 건너뛴다."""
    from claim_downloader.claims import kinds

    top = "4. 장학금 및 혁신인재지원금"
    sch_dir = _roster_folder(tmp_path, top, "근로장학금", "(나) TA", "9월", excel="2026학년도 (나)형 근로장학금 지급내역.xlsx")
    inn_dir = _roster_folder(tmp_path, top, "혁신인재지원금", "진행요원비", "9월")
    inn_named = _roster_folder(tmp_path, "a", "b", "c", "9월", excel="2026년도 9월 진행요원비 혁신인재지원금.xlsx")
    plain = _roster_folder(tmp_path, "x", "y", "z", "9월")
    assert kinds.guess(str(sch_dir)) == kinds.SCHOLARSHIP
    assert kinds.guess(str(inn_dir)) == kinds.INNOVATION          # 위 폴더 이름에서
    assert kinds.guess(str(inn_named)) == kinds.INNOVATION        # 명단 엑셀 이름에서
    assert kinds.guess(str(plain)) == kinds.SCHOLARSHIP           # 못 정하면 장학금


def test_innovation_job_uses_its_own_category_and_rules(monkeypatch, tmp_path):
    from claim_downloader import settings
    from claim_downloader.claims import kinds

    folder = _roster_folder(tmp_path, "혁신")
    job = scholarship.load_job(str(folder), kind=kinds.INNOVATION)
    assert job.spec.category == "교육·연구프로그램개발·운영비" and job.spec.code == "03"
    assert job.main_summary == "03_2026년도 7월 (가)형 근로장학금"            # 대표적요의 비목 번호만 03
    assert job.row_summary(0) == "03_2026년도 7월 (가)형 근로장학금_가나다"
    assert settings.innovation_rules_path().exists() and "# 혁신인재지원금 청구 적요 규칙" in settings.innovation_rules_path().read_text(encoding="utf-8")
    assert scholarship.load_job(str(folder)).main_summary.startswith("02_")     # 장학금은 그대로
    assert job.summary_lines()[0].startswith("종류: 혁신인재지원금")

    frame = FakeForm(reset_income=True)
    _patch_popups(monkeypatch, frame)
    logs = []
    report = scholarship.run(object(), frame, job, logs.append, ask=lambda m: False, wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and not report.stopped
    assert [row["PAY_EXP_CD"] for row in frame.saved] == ["교육·연구프로그램개발·운영비"] * 2
    assert all(row["INCM_DIV_CD"] == "장학금" for row in frame.saved)         # 소득구분은 장학금 그대로(공제 0)
    assert any("혁신인재지원금 청구 시작" in line for line in logs)
    assert any("✓ 비목 교육·연구프로그램개발·운영비 · 소득구분 장학금" in line for line in report.results[0].lines)


def test_total_row_with_other_labels_is_not_a_student_line(tmp_path):
    """'합계(원)' 처럼 적힌 합계 줄이 마지막 학생의 둘째 줄로 들어가 합계 금액이 청구될 뻔했다(2026-10-14 점검)."""
    path = make_na(tmp_path / "나.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws["B8"] = None
    ws["H8"] = "합계(원)"                        # 이름 칸 오른쪽에 적힌 합계
    wb.save(path)
    r = roster.read(path)
    assert [p.name for p in r.people] == ["사아자", "사아자", "차카타"] and r.sheet_total == 370000


def test_unmerged_blank_name_row_is_reported_not_inherited(tmp_path):
    path = make_na(tmp_path / "나.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws["J7"] = 999999                             # 셋째 학생 아래가 아니라 — 이름 없는 메모 줄
    ws.insert_rows(8)
    ws["J8"] = 5000
    wb.save(path)
    r = roster.read(path)
    assert all(p.amount != 5000 for p in r.people)
    assert any("이름이 없는데 금액" in w for w in r.warnings)


def test_korean_placeholders_in_nation_column_are_home():
    for value in ("-", "—", "한국인", "내국인", "대한민국(KR)", "해당없음", "KR"):
        assert roster.Person(name="가", nation=value).foreign_nation == "", value
    assert roster.Person(name="가", nation="UZ").foreign_nation == "UZ"


def test_float_student_number_does_not_gain_a_digit():
    assert roster.digits(202416354.0) == "202416354"


def test_owner_check_failure_alert_is_not_success(monkeypatch, tmp_path):
    """계좌를 넣고 [예금주조회] 했는데 ERP 가 실패 알림을 띄우고 결과 칸에 옛 '성공' 을 남기면 — 통과시키지 않는다(2026-10-14 점검)."""
    job = _mismatch_job(tmp_path, ["Account 110-594-395796"])
    frame = FakeForm(account="30120104410673", owner_check=None)   # 눌러도 결과 칸이 그대로
    _patch_popups(monkeypatch, frame)
    messages = []
    original = frame.evaluate

    def evaluate(js, arg=None):
        if js is payee.CLICK_ID_JS and arg == scholarship.OWNER_CHECK_ID:
            messages.append("예금주조회 실패: 계좌번호를 확인하세요")
        return original(js, arg)

    frame.evaluate = evaluate
    asked = []
    report = scholarship.run(object(), frame, job, lambda m: None, ask=lambda m: asked.append(m) or False,
                             wait=lambda s: None, messages=messages, dry_run=False)
    assert report.saved == 0 and asked and "실패" in asked[0]


def test_attachment_failure_is_not_uploaded_twice(monkeypatch, job):
    """앱이 첨부를 다 못 올리면 사람이 마저 올리고 [계속] — 앱은 다시 올리지 않고 첨부문서 수만 본다(전에는 겹쳐 올렸다)."""
    frame = FakeForm()
    calls = _patch_popups(monkeypatch, frame, attach_ok=False)

    def ask(message):
        frame.v["fileList"] = int(frame.v["fileList"]) + 2          # 사람이 ERP '첨부' 로 두 개를 올림
        return True

    report = scholarship.run(object(), frame, job, lambda m: None, ask=ask, wait=lambda s: None, dry_run=False)
    assert report.saved == 2 and len(calls["attach"]) == 2           # 학생마다 한 번씩만


# ── 신청 전 대조 (sheet_check.py) ─────────────────────────────────────────

def test_sheet_check_finds_duplicates_strangers_and_wrong_amounts(job):
    from claim_downloader.claims import sheet_check
    a, b = job.row_summary(0), job.row_summary(1)
    rows = [f"1 {a} 300,000", f"2 {b} 250,000", f"3 {a} 300,000", "4 02_2026년도 6월 (가)형 근로장학금_가나다 300,000"]
    check = sheet_check.compare(job, rows, [1150000, 1150000, 0])
    text = "\n".join(check.problems)
    assert "1. 가나다 — 청구내역에 2번 들어가 있습니다(줄 1, 3)" in text
    assert "2. 라마바 — 청구내역 줄 2 의 금액이 명단(240,000원)과 다릅니다" in text
    assert "명단에 없는 줄 4" in text
    assert not check.ok


def test_sheet_check_totals_deduction_and_missing(job):
    from claim_downloader.claims import sheet_check
    a = job.row_summary(0)
    check = sheet_check.compare(job, [f"1 {a} 300,000"], [300000, 273600, 26400])
    assert any("공제합계가 26,400원" in p for p in check.problems)
    assert any("청구내역에 없는 사람 1명: 라마바" in p for p in check.problems)
    # 넣다가 멈췄으면 빠진 사람은 당연하다 — 확인할 것이 아니라 줄로만
    partial = sheet_check.compare(job, [f"1 {a} 300,000"], [300000, 300000, 0], partial=True)
    assert partial.ok and any("라마바" in line for line in partial.lines)
    assert any("합계 줄 청구액 300,000원 = 명단" in line for line in partial.lines)


def test_sheet_check_does_not_list_every_row_when_nothing_matches(job):
    from claim_downloader.claims import sheet_check
    check = sheet_check.compare(job, ["1 03_소학회 진행요원비_가나다 300,000", "2 03_소학회 진행요원비_라마바 240,000"], [540000, 540000, 0])
    assert len(check.problems) == 1 and "하나도 맞추지 못했습니다" in check.problems[0]


def test_sheet_check_prefers_longest_summary(job):
    """'…_김민' 과 '…_김민수' 처럼 한 적요가 다른 적요를 품으면 긴 쪽에만 붙는다."""
    from claim_downloader.claims import sheet_check
    job.people[1].name = job.people[0].name + "수"
    a, b = job.row_summary(0), job.row_summary(1)
    check = sheet_check.compare(job, [f"1 {a} 300,000", f"2 {b} 240,000"], [540000, 540000, 0])
    assert check.ok, check.problems


def test_sheet_check_matches_truncated_summary_by_name(job):
    """ERP 청구내역 표는 긴 적요를 '…_ ...' 로 잘라 보인다(담당자 실제 화면 2026-10-19) — 대표적요 + 입금처 이름으로 맞춘다."""
    from claim_downloader.claims import sheet_check
    main = job.main_summary
    rows = [f"1 인건비(소득성 경비) 장학금 가나다 가나다 계좌이체 300,000 300,000 0 {main}_ ... [삭제]",
            f"2 인건비(소득성 경비) 장학금 라마바 라마바 계좌이체 240,000 240,000 0 {main}_ ... [삭제]"]
    check = sheet_check.compare(job, rows, [540000, 540000, 0])
    assert check.ok, check.problems
    assert any("잘려 보이는 줄 2개" in line for line in check.lines)
    # 같은 사람이 두 줄이면 여전히 '2번' 으로 잡는다
    twice = sheet_check.compare(job, rows + [rows[0].replace("1 ", "3 ", 1)], [840000, 840000, 0])
    assert any("1. 가나다 — 청구내역에 2번" in p for p in twice.problems)


def test_sheet_check_truncated_rows_of_another_claim_are_strangers(job):
    """이름이 같아도 대표적요가 다른 청구((나)형) · 너무 짧게 잘린 적요는 이 명단 줄로 보지 않는다."""
    from claim_downloader.claims import sheet_check
    a = job.row_summary(0)
    other = job.main_summary.replace("(가)형", "(나)형")
    rows = [f"1 가나다 300,000 {a}", f"2 라마바 240,000 {other}_ ...", "3 라마바 240,000 02_2026년도 7월 (가 ..."]
    check = sheet_check.compare(job, rows, [780000, 780000, 0])
    text = "\n".join(check.problems)
    assert "명단에 없는 줄 2" in text and "명단에 없는 줄 3" in text


def test_sheet_check_truncated_ta_rows_pick_by_amount(tmp_path):
    """(나)형 TA — 한 학생의 교과목 두 줄이 교과목 앞에서 잘리면 금액으로 명단 줄을 고른다."""
    from claim_downloader.claims import sheet_check
    folder = tmp_path / "TA"
    folder.mkdir()
    make_na(folder / "명단.xlsx")
    ta = scholarship.load_job(str(folder))
    main = ta.main_summary
    rows = [f"1 사아자 120,000 {main}(사아자_ ...", f"2 사아자 150,000 {main}(사아자_ ...", f"3 차카타 100,000 {main}(차카타_ ..."]
    check = sheet_check.compare(ta, rows, [370000, 370000, 0])
    assert check.ok, check.problems
    assert any("1. 사아자 — 줄 2" in line for line in check.lines)
    assert any("2. 사아자 — 줄 1" in line for line in check.lines)


# ── 중복 임시저장 알림 (record.py 이력) ───────────────────────────────────

def test_record_warning_names_people_already_saved_and_kind_change(isolated_home, job):
    from claim_downloader.claims import kinds
    from claim_downloader.claims import record
    assert scholarship.record_warning(job) == ""                     # 처음 넣는 폴더는 묻지 않는다
    report = scholarship.ScholarshipReport(dry_run=False)
    report.results = [scholarship.new_result(job, 0)]
    report.results[0].status = "saved"
    job.req_cnt = "324"
    scholarship.record_job(job, report)
    warning = scholarship.record_warning(job)
    assert "결의번호 324" in warning and "장학금" in warning and "1명" in warning
    assert "이번에 체크한 사람 중 1명은 이미 임시저장했습니다: 가나다" in warning
    # 같은 폴더를 다른 종류로 — 324 처럼 한 청구서에 비목이 섞이는 것을 미리 알린다
    other = scholarship.load_job(job.folder, kind=kinds.INNOVATION)
    assert "전에는 '장학금' 으로 넣었는데 지금은 '혁신인재지원금'" in scholarship.record_warning(other)
    # 두 번째로 넣으면 이력이 쌓이고(마지막 하나만 남던 옛 판과 달리) 저장한 사람도 합쳐진다
    report.results = [scholarship.new_result(job, 1)]
    report.results[0].status = "saved"
    scholarship.record_job(job, report)
    entry = record.previous(job.folder)
    assert len(entry["runs"]) == 2 and {s["name"] for s in entry["saved"]} == {"가나다", "라마바"}
    assert "2명은 이미 임시저장" in scholarship.record_warning(job)


def test_record_warning_finds_same_summary_in_another_folder(isolated_home, job, tmp_path):
    from claim_downloader.claims import kinds
    import shutil
    from claim_downloader.claims import record
    record.remember(job.folder, {"summary": job.main_summary + " (2명)", "req_cnt": "322"}, kind=kinds.SCHOLARSHIP)
    copy = tmp_path / "복사본"
    shutil.copytree(job.folder, copy)
    warning = scholarship.record_warning(scholarship.load_job(str(copy)))
    assert "다른 폴더에서 같은 대표적요로 넣은 적이 있습니다" in warning and "결의번호 322" in warning


def test_record_keeps_old_single_entry_as_history(isolated_home, tmp_path):
    """옛 판이 남긴 기록(마지막 하나뿐)도 이력의 첫 줄로 남는다."""
    import json
    from claim_downloader import settings
    from claim_downloader.claims import record
    folder = tmp_path / "옛"
    settings.claim_record_path().write_text(json.dumps({record._key(folder): {"when": "2026-10-13 10:00", "summary": "옛 것"}}),
                                            encoding="utf-8")
    record.remember(folder, {"summary": "새 것"})
    lines = record.history_lines(record.previous(folder))
    assert lines[0].startswith("2026-10-13 10:00") and len(lines) == 2


# ── 개별증빙 파일 이름이 소학회명인 경우 (담당자 실제 시험 2026-10-19) ──────────

def make_club_roster(path, rows):
    """진행요원비 명단 모양 — B2 제목, 3행 머리글 `연번 소학회명 이름 학번 은행 계좌 금액`. 이름 · 번호는 가짜."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "2026년도 9월 진행요원비(소학회 회장) 혁신인재지원금 지급내역"
    for c, h in enumerate(["연번", "소학회명", "이   름", "학   번", "은행", "계   좌", "금액"], start=2):
        ws.cell(3, c, h)
    for r, (club, name, no) in enumerate(rows, start=4):
        for c, v in enumerate([r - 3, club, name, no, "국민", "111-11-1111" + str(r), 240000], start=2):
            ws.cell(r, c, v)
    ws.cell(4 + len(rows), 2, "합   계")
    ws.cell(4 + len(rows), 8, 240000 * len(rows))
    wb.save(path)


def test_person_files_found_by_club_name(tmp_path):
    folder = tmp_path / "9월"
    (folder / "공통증빙").mkdir(parents=True)
    (folder / "개별증빙").mkdir()
    make_club_roster(folder / "진행요원비 지급내역.xlsx",
                     [("ALLERGY", "가나다", 202600001), ("BFS\n(BlockChain Fianacial Security)", "라마바", 202600002),
                      ("구름", "사아자", 202600003), ("구름자물쇠", "차카타", 202600004)])
    (folder / "공통증빙" / "공문.pdf").write_bytes(b"%PDF")
    names = ["AlLerGy_9월 근무상황부.pdf", "AlLerGy_9월 지원금 신청서.pdf", "BFS_9월 활동보고서.pdf",
             "구름_9월 근무상황부.pdf", "구름자물쇠_9월 근무상황부.pdf", "모르는팀_9월 근무상황부.pdf"]
    for n in names:
        (folder / "개별증빙" / n).write_bytes(b"%PDF")
    job = scholarship.load_job(str(folder))
    got = {job.people[i].group: sorted(os.path.basename(f) for f in files) for i, files in job.person_files.items()}
    assert got == {"ALLERGY": ["AlLerGy_9월 근무상황부.pdf", "AlLerGy_9월 지원금 신청서.pdf"],    # 대소문자는 가리지 않는다
                   "BFS\n(BlockChain Fianacial Security)": ["BFS_9월 활동보고서.pdf"],    # 칸이 두 줄이어도 첫 줄로
                   "구름": ["구름_9월 근무상황부.pdf"], "구름자물쇠": ["구름자물쇠_9월 근무상황부.pdf"]}  # 긴 이름이 먼저
    assert [os.path.basename(f) for f in job.unmatched] == ["모르는팀_9월 근무상황부.pdf"]
    assert any("어느 학생 것인지 모르는 개별증빙 1개" in p for p in job.problems())
    assert len(job.files_for(0)) == 3                                   # 공통 1 + 개별 2


def test_club_with_two_students_is_not_guessed(tmp_path):
    """한 소학회에 학생이 둘이면 소학회명만 든 파일은 누구 것인지 모른다 — 붙이지 않고 알린다. 이름이 든 파일은 그대로 붙는다."""
    folder = tmp_path / "9월"
    (folder / "개별증빙").mkdir(parents=True)
    make_club_roster(folder / "명단.xlsx", [("BFS", "가나다", 202600001), ("BFS", "라마바", 202600002)])
    for n in ("BFS_9월 활동보고서.pdf", "BFS_라마바_9월 신청서.pdf"):
        (folder / "개별증빙" / n).write_bytes(b"%PDF")
    job = scholarship.load_job(str(folder))
    assert {i: [os.path.basename(f) for f in v] for i, v in job.person_files.items()} == {1: ["BFS_라마바_9월 신청서.pdf"]}
    assert [os.path.basename(f) for f in job.unmatched] == ["BFS_9월 활동보고서.pdf"]
    # 개별증빙이 하나도 없는 사람은 미리 보기 '확인할 것' 에 보인다
    assert "개별증빙 없음" in scholarship.plan_person(job, 0).note
    assert "개별증빙 없음" not in (scholarship.plan_person(job, 1).note or "")
